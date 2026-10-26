from copy import deepcopy
from io import BytesIO
from uuid import uuid4
import pytest
from pptx import Presentation
from app.presentations.schemas import Slide, Outline, apply_selected
from app.presentations.rendering import catalog, validate_page, preview_svg, export_pptx


def page(layout='concept', title='光合作用'):
    slots = {s['id']: '学习内容' for s in catalog()[layout]['slots']}
    return {'id':str(uuid4()),'layout':layout,'title':title,'content':'观察叶片，理解植物如何获取能量。','slots':slots}


def test_selected_rewrite_never_accepts_extra_missing_or_duplicate_ids():
    a,b=page(),page()
    original=deepcopy([a,b]); changed={**a,'title':'叶片的秘密'}
    result=apply_selected([a,b],[a['id']],[changed])
    assert result[0]['title']=='叶片的秘密'
    assert result[1]==original[1] and [a,b]==original
    for bad in ([changed,b],[],[changed,changed]):
        with pytest.raises(ValueError):apply_selected([a,b],[a['id']],bad)


def test_outline_rejects_duplicate_ids_and_invalid_layout():
    p=page()
    with pytest.raises(ValueError):Outline(slides=[p,p])
    with pytest.raises(ValueError):Slide(**{**p,'layout':'arbitrary'})


def test_export_contains_all_pages_native_text_fixed_ratio_and_no_sample_labels():
    pages=[page(layout) for layout in catalog()]
    result=Presentation(BytesIO(export_pptx(pages,'植物的能量')))
    assert len(result.slides)==7
    assert abs(result.slide_width/result.slide_height-16/9)<0.001
    for slide in result.slides:
        text=' '.join(s.text for s in slide.shapes if s.has_text_frame)
        assert '学习内容' in text and '植物的能量' in text
        assert '二次函数' not in text and '{{' not in text


def test_preview_escapes_text_and_rejects_content_overflow_or_missing_slots():
    p=page(); key=next(iter(p['slots']));p['slots'][key]='<script>'
    svg=preview_svg(p,'课程',1)
    assert '<script>' not in svg and '&lt;script&gt;' in svg
    del p['slots'][key]
    with pytest.raises(ValueError):validate_page(p)
    p=page();p['slots'][key]='宽'*1000
    with pytest.raises(ValueError):validate_page(p)


def test_visual_slots_replace_entire_example_graph_and_native_shapes_have_no_theme_shadow():
    import xml.etree.ElementTree as ET
    p=page('image-text')
    p['slots']['teaching-visual']='阳光'+chr(10)+'植物生长'
    root=ET.fromstring(preview_svg(p,'植物',1))
    text=' '.join(n.text or '' for n in root.iter() if n.tag.endswith('text'))
    assert '(2, 4)' not in text and 'x' not in text and 'y' not in text
    assert '阳光' in text and '植物生长' in text
    deck=Presentation(BytesIO(export_pptx([p],'植物')))
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for shape in deck.slides[0].shapes:
        assert shape._element.find('.//a:effectLst',ns) is not None


def test_native_export_disables_theme_effects_and_automatic_line_wrapping():
    deck=Presentation(BytesIO(export_pptx([page('practice')],'植物')))
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for shape in deck.slides[0].shapes:
        for effect in shape._element.findall('.//a:effectRef',ns):
            assert effect.get('idx')=='0'
        if shape.shape_type==17:
            assert shape._element.find('.//a:bodyPr',ns).get('wrap')=='none'
            assert shape._element.find('.//a:noAutofit',ns) is not None


def test_chinese_wrap_keeps_closing_punctuation_with_text_and_opening_quotes_with_following_text():
    from app.presentations.rendering import wrap,units
    for value in ['叶片制造养分。','叶片制造。养分','观察“叶片颜色”的变化。','根吸水，叶片制造有机物。']:
        lines=wrap(value,4*28,28)
        assert ''.join(lines)==value
        assert all(not line.startswith(('。','，','、','；','：','！','？','”','）')) for line in lines)
        assert all(not line.endswith(('“','（')) for line in lines)
        assert all(units(line)*28<=4*28 for line in lines)
    assert units('“养分”')==4
