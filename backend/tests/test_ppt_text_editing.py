from copy import deepcopy
from io import BytesIO
from uuid import UUID,uuid4
import pytest
from pptx import Presentation
from app.presentations import text_editing as editor
from app.presentations.models import PptProject
from tests.test_ppt_api import create
from tests.test_ppt_master_runtime import sample_svg

def source():
    return {'id':str(uuid4()),'layout':'concept','title':'光合作用','content':'原正文','slots':{},'renderer':'ppt-master-6.7.0','svg':sample_svg(),'design':{'visual_brief':'叶片','composition':'图文并排'}}

def test_patch_changes_only_selected_text_and_styles():
    p=source();before=deepcopy(p)
    blocks=editor.text_blocks(p['svg'])
    assert blocks[0]['text']=='光合作用'
    result=editor.apply_text_edits(p,[{'key':blocks[0]['key'],'text':'植物的能量','font_size':40,'color':'#224466','bold':True}])
    assert p==before
    assert result['title']=='植物的能量'
    assert '植物的能量' in result['svg'] and '#224466' in result['svg']
    import xml.etree.ElementTree as ET
    assert ET.tostring(ET.fromstring(result['svg']).find('{http://www.w3.org/2000/svg}g'))==ET.tostring(ET.fromstring(p['svg']).find('{http://www.w3.org/2000/svg}g'))
    assert '原正文' not in result['content']

def test_literal_markup_is_text_and_overflow_is_rejected():
    p=source();key=editor.text_blocks(p['svg'])[0]['key']
    result=editor.apply_text_edits(p,[{'key':key,'text':'a < b & c > d'}])
    assert '&lt;' in result['svg'] and '<script' not in result['svg']
    with pytest.raises(editor.TextEditError):editor.apply_text_edits(p,[{'key':key,'text':'过长文字'*500}])
    with pytest.raises(editor.TextEditError):editor.apply_text_edits(p,[{'key':'not-found','text':'标题'}])

async def test_text_save_preserves_other_pages_history_and_export_gate(teacher_client,db):
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));first=source();second=source();second['title']='第二页'
    from app.presentations.delivery import certify_pages
    first,second=certify_pages([first,second],[b'old-first',b'old-second'])
    row.pages=[first,second];row.outline={'slides':[{k:x[k] for k in ('id','layout','title','content','slots')} for x in row.pages]};row.rendered_revision=row.confirmed_revision=1;row.status='ready';await db.commit()
    base='/api/presentations/'+p['id'];path=base+'/slides/'+first['id']
    descriptor=await teacher_client.get(path+'/editor?revision=1');assert descriptor.status_code==200,descriptor.text
    key=descriptor.json()['blocks'][0]['key']
    saved=await teacher_client.patch(path+'/text',json={'revision':1,'edits':[{'key':key,'text':'植物的能量'}]})
    assert saved.status_code==200,saved.text
    current=saved.json();assert current['revision']==2 and current['pages'][1]==second
    assert current['outline']['slides'][0]['title']=='植物的能量'
    assert current['versions'][0]['revision']==1
    assert not (await teacher_client.get('/api/presentations?scope=completed')).json()['items']
    old=await teacher_client.get(base+'/versions/1');assert old.json()['pages'][0]==first
    deck=await teacher_client.get(base+'/export');assert deck.status_code==200,deck.text
    assert '植物的能量' in ''.join(getattr(s,'text','') for s in Presentation(BytesIO(deck.content)).slides[0].shapes)
    stale=await teacher_client.patch(path+'/text',json={'revision':1,'edits':[{'key':key,'text':'旧修改'}]});assert stale.status_code==409
    assert (await teacher_client.get(path+'/editor?revision=1')).status_code==409

async def test_invalid_text_edit_keeps_last_deck(teacher_client,db):
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));page=source()
    row.pages=[page];row.outline={'slides':[{k:page[k] for k in ('id','layout','title','content','slots')}]};row.rendered_revision=1;row.status='ready';await db.commit()
    path='/api/presentations/'+p['id']
    bad=await teacher_client.patch(path+'/slides/'+page['id']+'/text',json={'revision':1,'edits':[{'key':'text-0','color':'url(javascript:bad)'}]})
    assert bad.status_code==422
    current=(await teacher_client.get(path)).json();assert current['revision']==1 and current['pages']==[page]


def test_unchanged_patch_is_noop_and_transformed_text_is_readonly():
    p=source();key=editor.text_blocks(p['svg'])[0]['key']
    assert editor.apply_text_edits(p,[{'key':key,'text':'光合作用'}])==p
    rotated=deepcopy(p);rotated['svg']=rotated['svg'].replace('<text x="64"','<text transform="rotate(4)" x="64"',1)
    assert not editor.text_blocks(rotated['svg'])[0]['editable']
    with pytest.raises(editor.TextEditError):editor.apply_text_edits(rotated,[{'key':key,'text':'标题'}])

async def test_other_owner_cannot_edit_and_invalid_data_preserves_history(teacher_client,db,make_user,make_client,app,monkeypatch):
    from app.presentations import router
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));page=source()
    row.pages=[page];row.outline={'slides':[{k:page[k] for k in ('id','layout','title','content','slots')}]};row.rendered_revision=1;row.status='ready';await db.commit()
    base='/api/presentations/'+p['id'];path=base+'/slides/'+page['id']
    await make_user('text-other');other=make_client(app);await other.post('/api/auth/login',json={'username':'text-other','password':'pass1234'})
    assert (await other.get(path+'/editor?revision=1')).status_code==404
    assert (await other.patch(path+'/text',json={'revision':1,'edits':[{'key':'text-0','text':'标题'}]})).status_code==404
    assert (await teacher_client.patch(path+'/text',json={'revision':1,'edits':[{'key':'text-0','color':'invalid'}]})).status_code==422
    saved=(await teacher_client.get(base)).json();assert saved['revision']==1 and not saved['versions'] and saved['pages']==[page]


def test_tspan_position_is_used_for_editable_multiline_text():
    p=source();p['svg']=p['svg'].replace('<text x="64" y="100" font-family="Microsoft YaHei, Arial" font-size="48" fill="#173359">光合作用</text>','<text font-family="Microsoft YaHei, Arial" font-size="48" fill="#173359"><tspan x="64" y="100">光合作用</tspan><tspan x="64" y="160">能量转化</tspan></text>')
    blocks=editor.text_blocks(p['svg']);assert blocks[0]['x']==64 and blocks[0]['y']==100
    result=editor.apply_text_edits(p,[{'key':blocks[0]['key'],'text':'植物的能量'}])
    assert '植物的能量' in result['svg']


def test_formula_runs_are_readonly_and_wrapped_heading_updates_metadata():
    p=source();original='<text x="64" y="100" font-family="Microsoft YaHei, Arial" font-size="48" fill="#173359">光合作用</text>'
    formula=deepcopy(p);formula['svg']=formula['svg'].replace(original,'<text x="64" y="100" font-size="28">H<tspan dy="10" font-size="18">2</tspan><tspan dy="-10">O</tspan></text>')
    b=editor.text_blocks(formula['svg'])[0];assert not b['editable']
    with pytest.raises(editor.TextEditError):editor.apply_text_edits(formula,[{'key':b['key'],'bold':True}])
    p['svg']=p['svg'].replace(original,'<text x="64" y="100" font-size="48" fill="#173359"><tspan x="64" y="100">光合</tspan><tspan x="64" y="165">作用</tspan></text>')
    changed=editor.apply_text_edits(p,[{'key':'text-0','text':'植物的能量'}]);assert changed['title']=='植物的能量'


def test_added_text_box_preserves_existing_nodes_and_persists_geometry():
    p=source();old=deepcopy(p)
    result=editor.apply_text_edits(p,[],additions=[{'key':'new-1','text':'新增课堂提问','x':64,'y':620,'width':400,'font_size':28,'color':'#224466','bold':True}])
    assert p==old and result['title']==p['title']
    import xml.etree.ElementTree as ET
    before=ET.fromstring(p['svg']);after=ET.fromstring(result['svg'])
    assert [ET.tostring(n) for n in before]==[ET.tostring(n) for n in list(after)[:-1]]
    box=editor.text_blocks(result['svg'])[-1]
    assert box['text']=='新增课堂提问' and box['x']==64 and box['y']==620 and box['width']==400
    assert '新增课堂提问' in result['content']
    with pytest.raises(editor.TextEditError):editor.apply_text_edits(p,[],additions=[{'key':'new-1','text':'超出页面'*100,'x':1200,'y':700,'width':400,'font_size':28}])

async def test_save_new_box_is_native_editable_text_and_history_is_untouched(teacher_client,db):
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));page=source()
    row.pages=[page];row.outline={'slides':[{k:page[k] for k in ('id','layout','title','content','slots')}]};row.rendered_revision=row.confirmed_revision=1;row.status='ready';await db.commit()
    base='/api/presentations/'+p['id'];path=base+'/slides/'+page['id']
    saved=await teacher_client.patch(path+'/text',json={'revision':1,'edits':[],'additions':[{'key':'new-1','text':'新增课堂提问','x':64,'y':620,'width':400,'font_size':28,'color':'#173359','bold':False}]})
    assert saved.status_code==200,saved.text
    assert saved.json()['revision']==2
    assert (await teacher_client.get(base+'/versions/1')).json()['pages'][0]==page
    exported=await teacher_client.get(base+'/export');assert exported.status_code==200,exported.text
    def native_text(shapes):
        return ''.join(getattr(shape,'text','') + (native_text(shape.shapes) if hasattr(shape,'shapes') else '') for shape in shapes)
    assert '新增课堂提问' in native_text(Presentation(BytesIO(exported.content)).slides[0].shapes)
    descriptor=await teacher_client.get(path+'/editor?revision=2');assert descriptor.status_code==200,descriptor.text
    box=descriptor.json()['blocks'][-1];assert box['text']=='新增课堂提问' and box['width']==400 and box['editable']
    changed=await teacher_client.patch(path+'/text',json={'revision':2,'edits':[{'key':box['key'],'text':'修改后的课堂提问'}]})
    assert changed.status_code==200,changed.text
    assert changed.json()['revision']==3
    exported=await teacher_client.get(base+'/export');assert exported.status_code==200,exported.text
    assert '修改后的课堂提问' in native_text(Presentation(BytesIO(exported.content)).slides[0].shapes)
    assert '新增课堂提问' in (await teacher_client.get(base+'/versions/2')).json()['pages'][0]['svg']



def test_reopened_added_box_can_edit_its_existing_footprint_near_other_text():
    p=source();p['svg']=p['svg'].replace('</svg>','<text x="64" y="280" font-size="28" fill="#173359">Nearby</text></svg>')
    added=editor.apply_text_edits(p,[],additions=[{'key':'new-1','text':'Added','x':64,'y':250,'width':400,'font_size':28}])
    key=editor.text_blocks(added['svg'])[-1]['key']
    edited=editor.apply_text_edits(added,[{'key':key,'text':'Changed'}])
    assert editor.text_blocks(edited['svg'])[-1]['text']=='Changed'


async def test_teacher_save_does_not_compile_render_or_acquire_export_capacity(teacher_client,db,monkeypatch):
    from app.presentations import router,quality
    from app.presentations.delivery import require_deliverable_pages
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));page=source()
    row.pages=[page];row.outline={'slides':[{k:page[k] for k in ('id','layout','title','content','slots')}]};row.rendered_revision=row.confirmed_revision=1;row.status='ready';await db.commit()
    def forbidden(*args,**kwargs):raise AssertionError('manual save must not run quality/compile/render')
    async def forbidden_async(*args,**kwargs):forbidden()
    with monkeypatch.context() as patch:
        patch.setattr(router,'compile_with_capacity_held',forbidden_async)
        patch.setattr(router,'require_export_capacity',forbidden_async)
        patch.setattr(quality,'render_review_images',forbidden)
        base='/api/presentations/'+p['id'];path=base+'/slides/'+page['id']
        saved=await teacher_client.patch(path+'/text',json={'revision':1,'edits':[{'key':'text-0','text':'教师确认的标题'}]})
        assert saved.status_code==200,saved.text
        state=saved.json();assert state['revision']==2
        assert 'review' not in state['pages'][0]
        assert state['pages'][0]['delivery']['mode']=='teacher_confirmed'
        require_deliverable_pages(state['pages'])
        changed=deepcopy(state['pages']);changed[0]['svg']=changed[0]['svg'].replace('#173359','#112233')
        assert changed[0]['svg']!=state['pages'][0]['svg']
        with pytest.raises(quality.QualityError):require_deliverable_pages(changed)
        assert (await teacher_client.get(path+'/editor?revision=2')).status_code==200
        assert (await teacher_client.get(path+'/preview')).status_code==200
        assert (await teacher_client.get(base+'/versions/1')).json()['pages'][0]==page
    with monkeypatch.context() as patch:
        original=router.master.run_tool
        def without_review(root,script,*args,**kwargs):
            if script=='svg_quality_checker.py':forbidden()
            return original(root,script,*args,**kwargs)
        patch.setattr(router.master,'run_tool',without_review)
        exported=await teacher_client.get(base+'/export');assert exported.status_code==200,exported.text
        assert '教师确认的标题' in ''.join(getattr(s,'text','') for s in Presentation(BytesIO(exported.content)).slides[0].shapes).replace('\n','').replace('\v','')

def test_teacher_text_overlap_is_not_a_save_blocker():
    p=source();p['svg']=p['svg'].replace('</svg>','<text x="64" y="170" font-size="28" fill="#173359">老师自行排版</text></svg>')
    changed=editor.apply_text_edits(p,[{'key':'text-0','text':'第一行标题\n第二行标题'}])
    assert '第二行标题' in changed['svg']


def test_mixed_deck_quality_gate_only_checks_generated_pages(monkeypatch):
    from app.presentations import master_runtime as master
    from app.presentations.delivery import confirm_teacher_page
    original=master.run_tool;checked=[]
    def record(root,script,*args,**kwargs):
        if script=='svg_quality_checker.py':
            checked.append(str(args[0]));return ''
        return original(root,script,*args,**kwargs)
    monkeypatch.setattr(master,'run_tool',record)
    generated=source();manual=confirm_teacher_page(source(),uuid4())
    with master.prepared_project([manual,generated],'混合课件',{}):pass
    assert len(checked)==1 and checked[0].endswith('P02.svg')
    checked.clear()
    with master.prepared_project([manual],'教师课件',{}):pass
    assert not checked
    checked.clear()
    with master.prepared_project([generated],'生成课件',{}):pass
    assert len(checked)==1 and not checked[0].endswith('.svg')


def test_generated_page_overlap_exception_uses_project_report_for_single_svg(tmp_path,monkeypatch):
    import json
    from types import SimpleNamespace
    from app.presentations import master_runtime as master
    project=tmp_path/'lesson';target=project/'svg_output'/'P02.svg'
    target.parent.mkdir(parents=True);target.write_text('<svg/>')
    validation=project/'validation';validation.mkdir()
    (validation/'svg_quality_report.json').write_text(json.dumps({'categories':{'blocking':{'issues':[{'message':'data-pptx-bounds overlaps another module'}]}}}))
    monkeypatch.setattr(master.subprocess,'run',lambda *args,**kwargs:SimpleNamespace(returncode=1,stdout='overlap report',stderr=''))
    assert master.run_tool(tmp_path,'svg_quality_checker.py',target)=='overlap report'


def test_text_width_patch_reflows_and_persists_without_changing_semantic_content():
    p=source();before=deepcopy(p)
    changed=editor.apply_text_edits(p,[{'key':'text-0','width':100}])
    block=editor.text_blocks(changed['svg'])[0]
    assert block['width']==100 and block['text']=='光合\n作用'
    assert changed['content']==p['content'] and changed['title']==p['title'] and p==before
    changed_again=editor.apply_text_edits(changed,[{'key':'text-0','width':400,'font_size':60}])
    assert editor.text_blocks(changed_again['svg'])[0]['width']==400
    assert editor.text_blocks(changed_again['svg'])[0]['font_size']==60
    for width in (0,20,1300,float('nan')):
        with pytest.raises(editor.TextEditError):editor.apply_text_edits(p,[{'key':'text-0','width':width}])

async def test_resize_save_reopen_move_and_native_font_export(teacher_client,db):
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));page=source()
    row.pages=[page];row.outline={'slides':[{k:page[k] for k in ('id','layout','title','content','slots')}]};row.rendered_revision=row.confirmed_revision=1;row.status='ready';await db.commit()
    base='/api/presentations/'+p['id'];path=base+'/slides/'+page['id']
    saved=await teacher_client.patch(path+'/text',json={'revision':1,'edits':[{'key':'text-0','width':360,'height':150,'font_size':48}]})
    assert saved.status_code==200,saved.text
    data=(await teacher_client.get(path+'/editor?revision=2')).json()
    assert data['blocks'][0]['width']==360 and data['blocks'][0]['font_size']==48 and data['blocks'][0]['height']==150
    obj=next(o for o in data['objects'] if o.get('text_key')=='text-0')
    moved=await teacher_client.patch(path+'/text',json={'revision':2,'edits':[{'key':'text-0','width':320,'height':120}],'moves':[{'key':obj['key'],'dx':30,'dy':20}]})
    assert moved.status_code==200,moved.text
    block=(await teacher_client.get(path+'/editor?revision=3')).json()['blocks'][0]
    assert block['width']==320 and block['x']==94 and block['y']==120 and block['font_size']==48 and block['height']==120
    exported=await teacher_client.get(base+'/export');assert exported.status_code==200,exported.text
    shape=next(s for s in Presentation(BytesIO(exported.content)).slides[0].shapes if '光合作用' in getattr(s,'text',''))
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt==36
    assert abs(shape.width/9525-320)<1 and abs(shape.height/9525-120)<1
    assert (await teacher_client.get(base+'/versions/1')).json()['pages'][0]==page


def test_text_frame_top_left_resize_and_move_keep_native_frame_in_sync():
    p=source();changed=editor.apply_text_edits(p,[{'key':'text-0','x':90,'y':130,'width':300,'height':140}])
    node=__import__('xml.etree.ElementTree',fromlist=['ElementTree']).fromstring(changed['svg'])[1]
    assert node.get('data-pptx-frame')=='90 82 300 140'
    obj=next(o for o in editor.movable_objects(changed['svg']) if o.get('text_key')=='text-0')
    moved=editor.apply_text_edits(changed,[],moves=[{'key':obj['key'],'dx':20,'dy':10}])
    block=editor.text_blocks(moved['svg'])[0]
    assert (block['x'],block['y'],block['width'],block['height'])==(110,140,300,140)
    node=__import__('xml.etree.ElementTree',fromlist=['ElementTree']).fromstring(moved['svg'])[1]
    assert node.get('data-pptx-frame')=='110 92 300 140'


def test_resizing_after_move_uses_final_slide_coordinates():
    p=source();obj=next(o for o in editor.movable_objects(p['svg']) if o.get('text_key')=='text-0')
    changed=editor.apply_text_edits(p,[{'key':'text-0','x':32,'y':100,'width':472,'height':70}],moves=[{'key':obj['key'],'dx':200,'dy':0}])
    block=editor.text_blocks(changed['svg'])[0]
    assert block['x']==32 and block['width']==472


def test_resized_automatic_line_breaks_do_not_enter_outline_on_other_text_edit():
    p=source();changed=editor.apply_text_edits(p,[{'key':'text-0','width':100}])
    changed=editor.apply_text_edits(changed,[{'key':'text-1','text':'新的正文'}])
    assert '光合作用' in changed['content'] and '光合\n作用' not in changed['content']
    assert changed['title']=='光合作用'

def test_added_box_frame_can_reach_canvas_edges_and_keeps_explicit_height():
    p=source();changed=editor.apply_text_edits(p,[],additions=[{'key':'new-1','text':'新文本','x':0,'y':28,'width':1200,'height':692,'font_size':28}])
    box=editor.text_blocks(changed['svg'])[-1]
    assert box['x']==0 and box['width']==1200 and box['height']==692

async def test_multiline_frame_dimensions_survive_native_export(teacher_client,db):
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));page=source()
    row.pages=[page];row.outline={'slides':[{k:page[k] for k in ('id','layout','title','content','slots')}]};row.rendered_revision=row.confirmed_revision=1;row.status='ready';await db.commit()
    base='/api/presentations/'+p['id'];path=base+'/slides/'+page['id']
    saved=await teacher_client.patch(path+'/text',json={'revision':1,'edits':[{'key':'text-0','width':100,'height':160}]})
    assert saved.status_code==200,saved.text
    descriptor=(await teacher_client.get(path+'/editor?revision=2')).json()
    assert descriptor['blocks'][0]['text']=='光合作用'
    exported=await teacher_client.get(base+'/export');assert exported.status_code==200,exported.text
    shape=next(s for s in Presentation(BytesIO(exported.content)).slides[0].shapes if '光合' in getattr(s,'text',''))
    assert abs(shape.width/9525-100)<1 and abs(shape.height/9525-160)<1


def test_shrink_then_move_to_right_edge_validates_final_frame():
    p=source();obj=next(o for o in editor.movable_objects(p['svg']) if o.get('text_key')=='text-0')
    changed=editor.apply_text_edits(p,[{'key':'text-0','width':80,'height':240}],moves=[{'key':obj['key'],'dx':1136,'dy':0}])
    block=editor.text_blocks(changed['svg'])[0]
    assert block['x']==1200 and block['width']==80
