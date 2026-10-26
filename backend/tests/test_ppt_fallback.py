from copy import deepcopy
from uuid import uuid4
import xml.etree.ElementTree as ET
import pytest
from app.presentations import fallback, master_runtime as master
from app.presentations.knowledge import VERSION, outline_digest, require_reviewed_outline

def outline(content, title='光合作用'):
    result={'slides':[{'id':str(uuid4()),'layout':'concept','title':title,'content':content,'slots':{}}]}
    result['review']={'version':VERSION,'status':'passed','issues':[],'changes':[],'content_sha256':outline_digest(result)}
    return result

def test_fallback_preserves_every_character_and_uses_readable_bounded_svg():
    source=outline(('x=-b/(2a)，a≠0；光能 → 化学能；<script> & 知识。'+chr(10))*60, '一'*80)
    pages=fallback.make_pages(source['slides'][0])
    assert len(pages)>1
    assert ''.join(p['content'] for p in pages)==source['slides'][0]['content']
    assert pages==fallback.make_pages(source['slides'][0])
    assert pages[0]['id']==source['slides'][0]['id']
    assert len({p['id'] for p in pages})==len(pages)
    for page in pages:
        master.validate_svg(page['svg'],{},require_visual=False)
        svg=ET.fromstring(page['svg'])
        assert not svg.findall('.//{http://www.w3.org/2000/svg}script')
        body=svg.find('.//{http://www.w3.org/2000/svg}g[@id="body"]')
        assert body is not None
        assert all(float(t.get('font-size'))>=28 for t in body)
        assert all(float(t.get('y'))<=624 for t in body)
        assert ''.join(t.text or '' for t in body)==page['content'].replace(chr(10),'')

def test_exact_pagination_preserves_knowledge_approval_and_other_slides():
    source=outline('正确内容。'*350)
    other=deepcopy(source['slides'][0]); other['id']=str(uuid4());other['content']='未修改的内容'
    source['slides'].append(other); source['review']['content_sha256']=outline_digest(source)
    original=deepcopy(source)
    pages=fallback.make_pages(source['slides'][0])
    derived=fallback.paginated_outline(source,{source['slides'][0]['id']:pages})
    require_reviewed_outline(derived)
    assert derived['slides'][-1]==other and source==original
    assert derived['review']['pagination']['source_sha256']==outline_digest(source)
    assert [s['id'] for s in derived['slides'][:-1]]==[p['id'] for p in pages]

def test_paginated_outline_rejects_altered_content_and_unreviewed_source():
    source=outline('知识点。'*200); sid=source['slides'][0]['id']
    pages=fallback.make_pages(source['slides'][0]);pages[-1]['content']+='错误'
    with pytest.raises(ValueError):fallback.paginated_outline(source,{sid:pages})
    source['slides'][0]['content']='未审查内容'
    with pytest.raises(ValueError):fallback.paginated_outline(source,{sid:fallback.make_pages(source['slides'][0])})

def test_page_limit_is_explicit_and_never_truncates():
    source=outline('内容。'*500)
    source['slides']=[{**source['slides'][0],'id':str(uuid4())} for _ in range(40)]
    source['review']['content_sha256']=outline_digest(source)
    sid=source['slides'][0]['id']
    with pytest.raises(ValueError,match='40'):fallback.paginated_outline(source,{sid:fallback.make_pages(source['slides'][0])})

from types import SimpleNamespace
from unittest.mock import AsyncMock
from app.presentations import generation
from app.presentations.quality import QualityError, audit_pages, require_approved_pages


@pytest.mark.parametrize('failure',['generation','review','planning'])
async def test_system_completes_failed_pages_with_reviewed_fixed_template(monkeypatch,failure):
    source=outline('光合作用利用光能。'*80)
    project=SimpleNamespace(outline=source,pages=[],title='自动完成')
    job=SimpleNamespace(kind='render',payload={'slide_ids':[],'instruction':''},checkpoint={})
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([],[])))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'load_assets',AsyncMock(return_value={}))
    sid=source['slides'][0]['id']
    design={'id':sid,'visual_kind':'text','visual_brief':'知识','composition':'宽松'}
    monkeypatch.setattr(generation,'plan_pages',AsyncMock(side_effect=generation.GenerationError('规划失败')) if failure=='planning' else AsyncMock(return_value={'pages':[design]}))
    from tests.test_ppt_master_runtime import sample_svg
    async def render_page(config,slide,*args,**kwargs):
        if failure=='generation':raise generation.GenerationError('图形无效')
        return {**slide,'svg':sample_svg(),'renderer':master.VERSION,'design':design}
    model=AsyncMock(side_effect=render_page);monkeypatch.setattr(generation,'render_page',model)
    monkeypatch.setattr(master,'compile_deck',lambda *args:b'PK-test')
    monkeypatch.setattr(master,'check_page',lambda *args:None)
    monkeypatch.setattr(generation,'render_review_images',lambda data,count:[b'exported-page']*count)
    async def review_page(config,page,image,context):
        assert image==b'exported-page'
        if not page.get('fallback'):return [{'category':'readability','message':'严重重叠','repair':'重新排版'}]
        assert any(s['id']==page['id'] for s in context['slides'])
        return []
    reviewer=AsyncMock(side_effect=review_page);monkeypatch.setattr(generation,'review_page',reviewer)
    async def progress(value,stage,checkpoint=None):
        if checkpoint is not None:job.checkpoint=deepcopy(checkpoint)
    options={'settings':SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA='),'progress':progress}
    result=await generation.build_result(None,project,job,**options)
    assert not result['partial'] and len(result['pages'])>1
    require_approved_pages(result['pages']); require_reviewed_outline(result['outline'])
    assert ''.join(p['content'] for p in result['pages'])==source['slides'][0]['content']
    assert all(p.get('fallback') for p in result['pages'])
    assert [p['id'] for p in result['pages']]==[s['id'] for s in result['outline']['slides']]
    reviewed=reviewer.await_count;generated=model.await_count
    resumed=await generation.build_result(None,project,job,**options)
    assert resumed==result and reviewer.await_count==reviewed and model.await_count==generated


@pytest.mark.parametrize('failure',['content','service'])
async def test_fallback_delivers_openable_file_without_falsely_passing_quality(monkeypatch,failure):
    source=outline('正确教学内容')
    project=SimpleNamespace(outline=source,pages=[],title='异常保留')
    job=SimpleNamespace(kind='render',payload={'slide_ids':[],'instruction':''},checkpoint={})
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([],[])))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'load_assets',AsyncMock(return_value={}))
    monkeypatch.setattr(generation,'plan_pages',AsyncMock(side_effect=generation.GenerationError('模型请求失败')))
    monkeypatch.setattr(master,'compile_deck',lambda *args:b'PK-test')
    monkeypatch.setattr(master,'check_page',lambda *args:None)
    monkeypatch.setattr(generation,'render_review_images',lambda data,count:[b'page']*count)
    reviewer=AsyncMock(side_effect=QualityError('服务暂不可用')) if failure=='service' else AsyncMock(return_value=[{'category':'teaching','message':'内容不一致','repair':'检查内容'}])
    monkeypatch.setattr(generation,'review_page',reviewer)
    async def progress(*args):pass
    result=await generation.build_result(None,project,job,SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA='),progress)
    from app.presentations.delivery import require_deliverable_pages
    require_deliverable_pages(result['pages'])
    assert not result['partial']
    assert all(p['review']['status']=='needs_attention' for p in result['pages'])
    assert all(p['delivery']['mode']=='fallback' for p in result['pages'])
    assert all(p.get('fallback') for p in result['pages'])
    assert reviewer.await_count>=1 and project.pages==[] and project.outline==source


async def test_review_only_does_not_consume_model_repair_budget():
    async def render(pages):return [b'exported']
    async def review(page,image):return [{'category':'readability','message':'不可读','repair':'修复'}]
    pages,_=await audit_pages([{'id':'x'}],selected_ids={'x'},render=render,review=review,repair=None,allow_partial=True)
    assert pages[0]['review']['repairs']==0
    assert pages[0]['review']['status']=='needs_attention'

def test_fixed_template_exports_as_editable_native_pptx():
    from io import BytesIO
    from pptx import Presentation
    from app.presentations.quality import render_review_images
    source=outline('光合作用：二氧化碳和水在光能作用下生成有机物并释放氧气。光能 → 化学能。x=-b/(2a)，a≠0。'*15)
    pages=fallback.make_pages(source['slides'][0])
    data=master.compile_deck(pages,'简洁模板验证',{})
    deck=Presentation(BytesIO(data))
    assert len(deck.slides)==len(pages)
    for slide,page in zip(deck.slides,pages):
        def text_shapes(shapes):
            for shape in shapes:
                if shape.has_text_frame:yield shape.text
                if hasattr(shape,'shapes'):yield from text_shapes(shape.shapes)
        texts=list(text_shapes(slide.shapes))
        assert ''.join(texts[1:-1])==page['content']
        assert not any(shape.shape_type==13 for shape in slide.shapes)
    images=render_review_images(data,len(pages))
    assert len(images)==len(pages) and all(image.startswith(bytes([137])+b'PNG') for image in images)


async def test_worker_restart_keeps_approved_fallback_parts(monkeypatch):
    source=outline('知识点。'*220)
    project=SimpleNamespace(outline=source,pages=[],title='中断恢复')
    job=SimpleNamespace(kind='render',payload={'slide_ids':[],'instruction':''},checkpoint={})
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([],[])))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'load_assets',AsyncMock(return_value={}))
    monkeypatch.setattr(generation,'plan_pages',AsyncMock(side_effect=generation.GenerationError('规划失败')))
    monkeypatch.setattr(master,'compile_deck',lambda *args:b'PK-test')
    monkeypatch.setattr(generation,'render_review_images',lambda data,count:[b'exported']*count)
    calls=[]
    async def review(config,page,image,context):
        calls.append(page['id'])
        if len(calls)==2:raise KeyboardInterrupt()
        return []
    monkeypatch.setattr(generation,'review_page',review)
    async def progress(value,stage,checkpoint=None):
        if checkpoint is not None:job.checkpoint=deepcopy(checkpoint)
    settings=SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA=')
    with pytest.raises(KeyboardInterrupt):await generation.build_result(None,project,job,settings,progress)
    first=deepcopy(job.checkpoint['fallback_groups'][source['slides'][0]['id']][0])
    assert first['review']['status']=='passed'
    result=await generation.build_result(None,project,job,settings,progress)
    require_approved_pages(result['pages'])
    assert result['pages'][0]==first and calls.count(first['id'])==1

@pytest.mark.parametrize('kind',['render','revise'])
async def test_worker_publishes_complete_fallback_continuations(teacher_client,db,settings,monkeypatch,kind):
    from sqlalchemy import update
    from app.presentations import router,tasks
    from app.presentations.models import PptProject
    from tests.test_ppt_api import create
    source=outline('已确认的教学知识。'*95)
    pages=fallback.make_pages(source['slides'][0])
    async def render(candidate):return [b'final']*len(candidate)
    async def review(page,image):return []
    pages,_=await audit_pages(pages,selected_ids={p['id'] for p in pages},render=render,review=review,repair=None)
    derived=fallback.paginated_outline(source,{source['slides'][0]['id']:pages})
    created=await create(teacher_client)
    from uuid import UUID
    values={'outline':source,'confirmed_revision':1}
    if kind=='revise':values.update(pages=[{**source['slides'][0],'renderer':master.VERSION}],rendered_revision=1,status='ready')
    await db.execute(update(PptProject).where(PptProject.id==UUID(created['id'])).values(**values))
    await db.commit()
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'build_result',AsyncMock(return_value={'outline':derived,'pages':pages,'warnings':[],'partial':False}))
    payload={'revision':1,'idempotency_key':'fallback-'+kind}
    if kind=='revise':payload.update(slide_ids=[source['slides'][0]['id']],instruction='改进展示')
    base='/api/presentations/'+created['id']
    response=await teacher_client.post(base+'/jobs/'+kind,json=payload)
    assert response.status_code==200,response.text
    await tasks.run_job(response.json()['id'],settings)
    state=(await teacher_client.get(base)).json()
    assert state['job']['status']=='succeeded' and state['status']=='ready'
    assert state['pages']==pages and state['outline']==derived
    assert state['confirmed_revision']==state['rendered_revision']==state['revision']
    assert (await teacher_client.get(base+'/slides/'+pages[-1]['id']+'/preview')).status_code==200

@pytest.mark.parametrize('content',[chr(10)*120+'知识点'+chr(10)*120,'知识点'+chr(10)*120+'下一点','公式'+chr(9)*60+'条件', '第一点'+chr(13)+'第二点'])
def test_whitespace_heavy_source_remains_exact_without_blank_continuations(content):
    source=outline(content)
    pages=fallback.make_pages(source['slides'][0])
    assert ''.join(page['content'] for page in pages)==content
    assert all(page['content'].strip() for page in pages)
    fallback.paginated_outline(source,{source['slides'][0]['id']:pages})


async def test_selected_revision_fallback_keeps_other_actual_pages(monkeypatch):
    source=outline('保留页')
    second=deepcopy(source['slides'][0]); second.update(id=str(uuid4()),title='需修改页',content='教学要点。'*150)
    source['slides'].append(second);source['review']['content_sha256']=outline_digest(source)
    good=fallback.make_pages(source['slides'][0])[0]
    async def render(candidate):return [b'good']
    async def review(page,image):return []
    checked,_=await audit_pages([good],selected_ids={good['id']},render=render,review=review,repair=None)
    project=SimpleNamespace(outline=source,pages=[checked[0],{**second,'renderer':master.VERSION}],title='选页修改')
    job=SimpleNamespace(kind='revise',payload={'slide_ids':[second['id']],'instruction':'优化展示'},checkpoint={})
    monkeypatch.setattr(generation,'review_revision_group',AsyncMock(return_value=[]))
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([],[])))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'load_assets',AsyncMock(return_value={}))
    monkeypatch.setattr(generation,'rewrite_outline',AsyncMock(return_value=source))
    import app.presentations.knowledge as knowledge
    monkeypatch.setattr(knowledge,'check_outline',AsyncMock(return_value=source))
    monkeypatch.setattr(generation,'plan_pages',AsyncMock(side_effect=generation.GenerationError('规划失败')))
    monkeypatch.setattr(master,'compile_deck',lambda *args:b'PK-test')
    monkeypatch.setattr(generation,'render_review_images',lambda data,count:[b'exported']*count)
    monkeypatch.setattr(generation,'review_page',AsyncMock(return_value=[]))
    async def progress(*args):pass
    result=await generation.build_result(None,project,job,SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA='),progress)
    assert result['pages'][0]==checked[0] and len(result['pages'])>2
    assert all(p['fallback']['source_slide_id']==second['id'] for p in result['pages'][1:])
    require_approved_pages(result['pages'])

@pytest.mark.parametrize('separator',[chr(0x85),chr(0x2028),chr(0x2029)])
def test_copied_unicode_line_breaks_preserve_exact_content(separator):
    source=outline('知识点'+separator+'适用条件')
    pages=fallback.make_pages(source['slides'][0])
    assert ''.join(page['content'] for page in pages)==source['slides'][0]['content']
    fallback.paginated_outline(source,{source['slides'][0]['id']:pages})


def test_title_with_many_line_breaks_still_has_bounded_geometry():
    source=outline('教学内容',('短标题'+chr(10))*10)
    pages=fallback.make_pages(source['slides'][0])
    master.validate_svg(pages[0]['svg'],{},require_visual=False)
    assert pages[0]['title']==source['slides'][0]['title']
    assert max(float(t.get('y')) for t in ET.fromstring(pages[0]['svg']).findall('.//{http://www.w3.org/2000/svg}text'))<=684


def test_native_fallback_preserves_spacing_in_text():
    from io import BytesIO
    from pptx import Presentation
    source=outline(' a  b    c ')
    pages=fallback.make_pages(source['slides'][0])
    deck=Presentation(BytesIO(master.compile_deck(pages,'空白保留',{})))
    def text_shapes(shapes):
        for shape in shapes:
            if shape.has_text_frame:yield shape.text
            if hasattr(shape,'shapes'):yield from text_shapes(shape.shapes)
    texts=list(text_shapes(deck.slides[0].shapes))
    assert ''.join(texts[1:-1])==source['slides'][0]['content']
