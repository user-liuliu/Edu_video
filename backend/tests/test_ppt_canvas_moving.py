from copy import deepcopy
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile
import xml.etree.ElementTree as ET
import pytest
from app.presentations import text_editing as editor
from app.presentations.models import PptProject
from tests.test_ppt_text_editing import source
from tests.test_ppt_api import create

NS='{http://www.w3.org/2000/svg}'

def objects(page):
    return editor.movable_objects(page['svg'])

def test_descriptors_include_groups_shapes_and_text_but_not_background():
    rows=objects(source())
    assert {'group','text','shape'}<={r['kind'] for r in rows}
    assert all(r['path']!=[0] for r in rows)
    assert all(len(r['bounds'])==4 and r['bounds'][2]>=0 for r in rows)

def test_move_group_preserves_path_commands_and_text_stays_editable():
    p=source();old=deepcopy(p);group=next(r for r in objects(p) if r['kind']=='group')
    moved=editor.apply_text_edits(p,[],moves=[{'key':group['key'],'dx':40,'dy':-20}])
    assert p==old and moved['content']==p['content']
    before=ET.fromstring(p['svg']);after=ET.fromstring(moved['svg'])
    assert ET.tostring(before[0])==ET.tostring(after[0])
    assert after.find(NS+'g').get('data-pptx-bounds')=='520 140 650 480'
    rows=editor.text_blocks(moved['svg']);assert rows[-1]['x']==720 and rows[-1]['y']==560 and rows[-1]['editable']
    from app.presentations.master_runtime import validate_svg
    validate_svg(moved['svg'],{})
    changed=editor.apply_text_edits(moved,[{'key':rows[-1]['key'],'text':'新的叶片说明'}])
    assert '新的叶片说明' in changed['svg']

def test_text_and_image_moves_bake_coordinates_and_reject_outside_canvas():
    p=source();p['svg']=p['svg'].replace('</svg>','<image x="80" y="200" width="320" height="180" href="../images/aaaaaaaaaaaaaaaaaaaaaaaa.png"/></svg>')
    rows=objects(p);image=next(r for r in rows if r['kind']=='image');text=next(r for r in rows if r.get('text_key')=='text-0')
    changed=editor.apply_text_edits(p,[],moves=[{'key':image['key'],'dx':100,'dy':20},{'key':text['key'],'dx':20,'dy':10}])
    root=ET.fromstring(changed['svg']);assert root.find(NS+'image').get('x')=='180'
    assert editor.text_blocks(changed['svg'])[0]['x']==84
    assert editor.apply_text_edits(p,[],moves=[{'key':image['key'],'dx':0,'dy':0}])==p
    for move in [{'key':image['key'],'dx':2000,'dy':0},{'key':'object-9999','dx':2,'dy':0},{'key':image['key'],'dx':float('nan'),'dy':0}]:
        with pytest.raises(editor.TextEditError):editor.apply_text_edits(p,[],moves=[move])

def test_move_inside_existing_transform_does_not_introduce_nested_transforms():
    p=source();p['svg']=p['svg'].replace('</svg>','<g transform="translate(100 80) scale(2)"><rect x="20" y="30" width="80" height="40" fill="#224466"/></g></svg>')
    shape=next(r for r in objects(p) if r['path']==[3,0])
    moved=editor.apply_text_edits(p,[],moves=[{'key':shape['key'],'dx':40,'dy':20}])
    root=ET.fromstring(moved['svg']);rect=root[-1][0]
    assert float(rect.get('x'))==40 and float(rect.get('y'))==40 and not rect.get('transform')
    from app.presentations.master_runtime import validate_svg
    validate_svg(moved['svg'],{})

async def test_move_save_preserves_history_other_pages_and_exports_native_positions(teacher_client,db):
    p=await create(teacher_client);row=await db.get(PptProject,UUID(p['id']));page=source();other=source()
    from app.presentations.delivery import certify_pages
    page,other=certify_pages([page,other],[b'old',b'other'])
    row.pages=[page,other];row.outline={'slides':[{k:x[k] for k in ('id','layout','title','content','slots')} for x in row.pages]};row.rendered_revision=row.confirmed_revision=1;row.status='ready';await db.commit()
    base='/api/presentations/'+p['id'];path=base+'/slides/'+page['id']
    descriptor=await teacher_client.get(path+'/editor?revision=1');assert descriptor.status_code==200
    group=next(r for r in descriptor.json()['objects'] if r['kind']=='group')
    saved=await teacher_client.patch(path+'/text',json={'revision':1,'moves':[{'key':group['key'],'dx':40,'dy':-20}]})
    assert saved.status_code==200,saved.text
    assert saved.json()['revision']==2 and saved.json()['pages'][1]==other
    assert (await teacher_client.get(base+'/versions/1')).json()['pages'][0]==page
    reopened=await teacher_client.get(path+'/editor?revision=2');assert reopened.status_code==200,reopened.text
    assert next(r for r in reopened.json()['objects'] if r['key']==group['key'])['bounds'][0]==group['bounds'][0]+40
    exported=await teacher_client.get(base+'/export');assert exported.status_code==200,exported.text
    from pptx import Presentation
    deck=Presentation(BytesIO(exported.content));moved_group=next(s for s in deck.slides[0].shapes if hasattr(s,'shapes'))
    from app.presentations.master_runtime import compile_deck
    original_deck=Presentation(BytesIO(compile_deck([page],p['title'],{})))
    original_group=next(s for s in original_deck.slides[0].shapes if hasattr(s,'shapes'))
    assert abs((moved_group.left-original_group.left)/9525-40)<1
    assert abs((moved_group.top-original_group.top)/9525+20)<1
    with ZipFile(BytesIO(exported.content)) as z:assert 'custGeom' in z.read('ppt/slides/slide1.xml').decode()
    invalid=await teacher_client.patch(path+'/text',json={'revision':2,'moves':[{'key':group['key'],'dx':3000,'dy':0}]})
    assert invalid.status_code==422
    assert (await teacher_client.get(base)).json()['revision']==2


def test_group_and_child_moves_are_order_independent_in_world_coordinates():
    p=source();rows=objects(p);group=next(r for r in rows if r['kind']=='group');text=next(r for r in rows if r.get('text_key')=='text-1')
    patches=[{'key':group['key'],'dx':30,'dy':-20},{'key':text['key'],'dx':10,'dy':5}]
    a=editor.apply_text_edits(p,[],moves=patches)
    b=editor.apply_text_edits(p,[],moves=list(reversed(patches)))
    assert editor.text_blocks(a['svg'])[-1]['x']==720
    assert editor.text_blocks(a['svg'])[-1]['y']==565
    assert a['svg']==b['svg']


def test_transformed_group_can_move_left_while_staying_inside_canvas():
    p=source();p['svg']=p['svg'].replace('</svg>','<g data-pptx-bounds="20 30 80 40" transform="translate(200 100) scale(2)"><rect x="20" y="30" width="80" height="40" fill="#224466"/></g></svg>')
    group=next(r for r in objects(p) if r['path']==[3])
    moved=editor.apply_text_edits(p,[],moves=[{'key':group['key'],'dx':-180,'dy':0}])
    from app.presentations.master_runtime import validate_svg
    validate_svg(moved['svg'],{})
    assert next(r for r in objects(moved) if r['key']==group['key'])['bounds'][0]==60

def test_text_with_default_x_moves_and_tspan_inheritance_is_not_moved_twice():
    p=source();p['svg']=p['svg'].replace('</svg>','<text y="100" font-size="28" fill="#173359">Title<tspan> tail</tspan></text></svg>')
    item=next(r for r in objects(p) if r['path']==[3])
    moved=editor.apply_text_edits(p,[],moves=[{'key':item['key'],'dx':100,'dy':20}])
    node=ET.fromstring(moved['svg'])[-1]
    assert node.get('x')=='100' and node.get('y')=='120'
    assert node[0].get('x') is None and node[0].get('y') is None


def test_page_background_group_is_fixed_with_its_decorations():
    p=source();p['svg']=p['svg'].replace('<rect id="background"','<g id="page-background"><rect id="background"').replace('fill="#FFFFFF"/>','fill="#FFFFFF"/><circle cx="10" cy="10" r="5" fill="#224466"/></g>',1)
    assert not any(r['path'][0]==0 for r in objects(p))


def test_moved_picture_remains_native_image_with_exported_position():
    from app.presentations.master_runtime import compile_deck
    from tests.test_ppt_master_assets import png
    from pptx import Presentation
    p=source();name='a'*24+'.png'
    p['svg']=p['svg'].replace('</svg>',f'<image x="80" y="200" width="320" height="180" href="../images/{name}"/></svg>')
    image=next(r for r in objects(p) if r['kind']=='image')
    moved=editor.apply_text_edits(p,[],moves=[{'key':image['key'],'dx':100,'dy':20}])
    deck=Presentation(BytesIO(compile_deck([moved],'图片移动',{name:png()})))
    picture=next(s for s in deck.slides[0].shapes if s.shape_type==13)
    assert abs(picture.left/9525-180)<1 and abs(picture.top/9525-220)<1
    assert abs(picture.width/9525-320)<1 and abs(picture.height/9525-180)<1
