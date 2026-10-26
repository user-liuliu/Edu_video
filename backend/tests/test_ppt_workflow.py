from datetime import datetime,UTC,timedelta
from io import BytesIO
from unittest.mock import AsyncMock
from uuid import UUID,uuid4
import json
import pytest
from pptx import Presentation
from sqlalchemy import select,update
from app.presentations import generation,tasks
from app.presentations.models import PptJob,PptProject
from tests.test_ppt_api import create
from app.presentations.knowledge import outline_digest,VERSION as KNOWLEDGE_VERSION

def checked_outline(outline):
    return {**outline,'review':{'version':KNOWLEDGE_VERSION,'status':'passed','issues':[],'changes':[],'content_sha256':outline_digest(outline)}}

async def test_ppt_full_workflow_and_page_revision_preserve_unselected(teacher_client,db,settings,monkeypatch):
    from app.presentations import router
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    async def model(config,messages,**kwargs):
        content=messages[-1]['content']
        request=json.loads(content[0]['text'] if isinstance(content,list) else content)
        if request['operation']=='knowledge_review':return {'status':'passed','issues':[]}
        if request['operation'] in ('review','revision_review'):return {'status':'passed','issues':[]}
        if request['operation']=='outline':return {'slides':[{'layout':'concept','title':'光合作用','content':'植物使用阳光制造养分'},{'layout':'summary','title':'总结','content':'观察植物与环境'}]}
        if request['operation']=='edit_svg':
            from tests.test_ppt_master_runtime import sample_svg
            return {'svg':sample_svg().replace('光合作用','修改后的标题'),'slide':{**request['slide'],'title':'修改后的标题','content':'加入叶绿体实验'}}
        if request['operation']=='rewrite':
            return {'slides':[{**s,'title':'修改后的标题','content':'加入叶绿体实验'} for s in request['slides'] if s['id'] in request['selected_ids']]}
        if request['operation']=='design':
            assert 'OFFICIAL strategist.md' in messages[0]['content']
            return {'pages':[{'id':s['id'],'visual_kind':'diagram','visual_brief':'叶片、阳光和物质流动','composition':'图文并排'} for s in request['slides']]}
        if request['operation']=='svg':
            from tests.test_ppt_master_runtime import sample_svg
            assert 'OFFICIAL executor-base.md' in messages[0]['content']
            return {'svg':sample_svg().replace('光合作用',request['slide']['title'])}
        raise AssertionError(request['operation'])
    monkeypatch.setattr(generation,'chat_json',model)
    p=await create(teacher_client);base='/api/presentations/'+p['id']
    async def start(kind,p,**extra):
        r=await teacher_client.post(base+'/jobs/'+kind,json={'revision':p['revision'],'idempotency_key':str(uuid4()),**extra});assert r.status_code==200,r.text
        await tasks.run_job(r.json()['id'],settings)
        state=(await teacher_client.get(base)).json()
        assert state['job']['status']=='succeeded',state
        return state
    p=await start('outline',p);assert len(p['outline']['slides'])==2
    saved=await teacher_client.put(base+'/outline',json={'revision':p['revision'],'slides':p['outline']['slides']});assert saved.status_code==200
    confirmed=await teacher_client.post(base+'/confirm',json={'revision':saved.json()['revision']})
    assert confirmed.status_code==200,confirmed.text
    assert confirmed.json()['job']['kind']=='outline'  # Confirmation creates no generation job.
    p=await start('render',confirmed.json());assert len(p['pages'])==2
    assert all(page['review']['status']=='passed' for page in p['pages'])
    assert all(page['renderer']=='ppt-master-6.7.0' and 'teaching-visual' in page['svg'] for page in p['pages'])
    unchanged=p['pages'][1]
    p=await start('revise',p,slide_ids=[p['pages'][0]['id']],instruction='加入实验')
    assert p['pages'][1]==unchanged and p['pages'][0]['title']=='修改后的标题'
    r=await teacher_client.get(base+'/export');assert r.status_code==200
    deck=Presentation(BytesIO(r.content));assert len(deck.slides)==2
    preview=await teacher_client.get(base+'/slides/'+unchanged['id']+'/preview')
    assert preview.status_code==200 and 'teaching-visual' in preview.text
    from zipfile import ZipFile
    with ZipFile(BytesIO(r.content)) as z:
        assert 'custGeom' in z.read('ppt/slides/slide1.xml').decode()

async def test_stale_worker_cannot_publish_or_complete_reclaimed_job(teacher_client,db,settings,monkeypatch):
    from app.presentations import router
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    p=await create(teacher_client);base='/api/presentations/'+p['id']
    r=await teacher_client.post(base+'/jobs/outline',json={'revision':1,'idempotency_key':'stale'});jid=UUID(r.json()['id'])
    replacement=uuid4()
    async def steal(*args,**kwargs):
        await db.execute(update(PptJob).where(PptJob.id==jid).values(lease_token=replacement,lease_until=datetime.now(UTC)+timedelta(seconds=90)))
        await db.commit()
        return {'outline':{'slides':[]},'pages':[],'warnings':[]}
    monkeypatch.setattr(generation,'build_result',steal)
    await tasks.run_job(jid,settings)
    project=await db.get(PptProject,UUID(p['id']),populate_existing=True);job=await db.get(PptJob,jid,populate_existing=True)
    assert project.outline is None and job.lease_token==replacement and job.status=='running'

async def test_outbox_recovers_expired_leases_and_broker_failure_keeps_pending(teacher_client,db,monkeypatch):
    from app.presentations import router
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    p=await create(teacher_client)
    r=await teacher_client.post('/api/presentations/'+p['id']+'/jobs/outline',json={'revision':1,'idempotency_key':'outbox'})
    jid=UUID(r.json()['id']);calls=[]
    def broken(value):raise ConnectionError('broker unavailable')
    await tasks.dispatch_pending(db,broken)
    job=await db.get(PptJob,jid,populate_existing=True);assert job.dispatched_at is None and job.status=='pending'
    job.status='running';job.lease_until=datetime.now(UTC)-timedelta(seconds=10);await db.commit()
    await tasks.dispatch_pending(db,calls.append)
    assert calls==[str(jid)]

async def test_failed_reference_identifies_file_and_preserves_actionable_reason(teacher_client,db,settings,monkeypatch):
    from app.presentations.models import PptMaterial
    from app.presentations.generation import references,GenerationError
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    material=PptMaterial(owner_id=project.owner_id,project_id=project.id,filename='课堂图片.png',object_key='test-file',content_type='image/png',size=100)
    db.add(material);await db.commit();mid=material.id
    from types import SimpleNamespace
    monkeypatch.setattr(generation,'MaterialStorage',lambda settings:SimpleNamespace(get=lambda key:b'image'))
    monkeypatch.setattr(generation,'parse_document',lambda *a,**kw:[{'location':'图片 1','text':'','image':b'data'}])
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=None))
    import pytest
    with pytest.raises(GenerationError):await references(db,project,settings)
    await db.refresh(material)
    assert material.status=='failed'
    assert '看图' in material.error_message and '课堂图片.png' in material.error_message
    retry=await teacher_client.post('/api/presentations/'+p['id']+'/materials/'+str(mid)+'/retry?revision=1')
    assert retry.status_code==200 and retry.json()['status']=='stored'

async def test_old_ready_material_reparsed_and_image_retained_with_provenance(teacher_client,db,settings,monkeypatch):
    from app.presentations.models import PptMaterial
    from app.presentations.master_assets import load_assets
    from tests.test_ppt_master_assets import png
    from types import SimpleNamespace
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    material=PptMaterial(owner_id=project.owner_id,project_id=project.id,filename='leaf.png',object_key='owned-original',content_type='image/png',size=500,status='ready',chunks=[{'location':'图片1','text':'旧识别结果'}])
    db.add(material);await db.commit()
    blob=png()
    monkeypatch.setattr(generation,'MaterialStorage',lambda s:SimpleNamespace(get=lambda key:blob))
    monkeypatch.setattr(generation,'parse_document',lambda *a,**k:[{'location':'图片1','text':'','image':blob}])
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'recognize_image',AsyncMock(return_value='叶片参考图'))
    from app.presentations import master_assets
    monkeypatch.setattr(master_assets,'MaterialStorage',lambda s:SimpleNamespace(get=lambda key:blob))
    monkeypatch.setattr(master_assets,'parse_document',lambda *a,**k:[{'location':'图片1','text':'','image':blob}])
    refs,_=await generation.references(db,project,settings)
    assert refs[0]['image']['name'].endswith('.png')
    assert material.chunks[0]['parse_version']==2 and 'image' in material.chunks[0]
    assets=await load_assets(db,project,settings)
    assert list(assets)==[refs[0]['image']['name']]
    assert assets[refs[0]['image']['name']][:8]==bytes([137,80,78,71,13,10,26,10])
    await generation.references(db,project,settings)
    assert generation.recognize_image.await_count==1


async def test_failed_master_compile_preserves_published_deck(teacher_client,db,settings,monkeypatch):
    from app.presentations import router,master_runtime
    from tests.test_ppt_master_runtime import sample_svg
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    sid=str(uuid4())
    slide={'id':sid,'layout':'image-text','title':'原有标题','content':'叶片与阳光','slots':{}}
    page={**slide,'renderer':'ppt-master-6.7.0','svg':sample_svg(),'design':{'id':sid,'visual_kind':'diagram','visual_brief':'叶片和阳光','composition':'左文右图'}}
    project.outline=checked_outline({'slides':[slide]});project.pages=[page];project.rendered_revision=1;project.confirmed_revision=1;project.status='ready'
    await db.commit()
    async def model(config,messages,**kwargs):
        req=json.loads(messages[-1]['content'])
        if req['operation']=='knowledge_review':return {'status':'passed','issues':[]}
        if req['operation']=='rewrite':return {'slides':[{**slide,'title':'新标题'}]}
        if req['operation']=='design':return {'pages':[page['design']]}
        if req['operation']=='svg':return {'svg':sample_svg().replace('光合作用','新标题')}
        raise AssertionError(req['operation'])
    monkeypatch.setattr(generation,'chat_json',model)
    def broken(*args,**kwargs):raise master_runtime.MasterError('页面超出边界，请调整布局后重试')
    monkeypatch.setattr(master_runtime,'compile_deck',broken)
    base='/api/presentations/'+p['id']
    r=await teacher_client.post(base+'/jobs/revise',json={'revision':1,'idempotency_key':'compile-failure','slide_ids':[sid],'instruction':'修改标题'})
    assert r.status_code==200,r.text
    await tasks.run_job(r.json()['id'],settings)
    state=(await teacher_client.get(base)).json()
    assert state['job']['status']=='failed' and '原课件已保留' in state['job']['error_message']
    assert state['pages']==[page] and state['revision']==1 and state['rendered_revision']==1


async def test_reference_preview_and_export_are_owned_embedded_and_invalidated(teacher_client,db,settings,monkeypatch,make_user,make_client,app):
    from app.presentations import master_assets
    from app.presentations.models import PptMaterial
    from app.presentations.master_assets import normalize_image,image_descriptor
    from tests.test_ppt_master_assets import png
    from tests.test_ppt_master_runtime import sample_svg
    from types import SimpleNamespace
    from zipfile import ZipFile
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    raw=png();data=normalize_image(raw);mid=uuid4();desc=image_descriptor(mid,0,data);sid=str(uuid4())
    material=PptMaterial(id=mid,owner_id=project.owner_id,project_id=project.id,filename='leaf.png',object_key='private-original',content_type='image/png',size=len(raw),status='ready',chunks=[{'location':'图片1','text':'叶片','parse_version':2,'image':desc}])
    db.add(material)
    slide={'id':sid,'layout':'image-text','title':'叶片','content':'观察叶片','slots':{}}
    svg=sample_svg().replace('</g></svg>',f'<image href="../images/{desc["name"]}" x="80" y="200" width="320" height="180"/></g></svg>')
    page={**slide,'renderer':'ppt-master-6.7.0','svg':svg,'design':{'id':sid,'visual_kind':'reference-image','visual_brief':'叶片图解与参考图片','composition':'图文并排'}}
    project.outline=checked_outline({'slides':[slide]});project.pages=[page];project.rendered_revision=1;project.status='ready';await db.commit()
    monkeypatch.setattr(master_assets,'MaterialStorage',lambda s:SimpleNamespace(get=lambda key:raw))
    monkeypatch.setattr(master_assets,'parse_document',lambda *a,**k:[{'location':'图片1','text':'','image':raw}])
    base='/api/presentations/'+p['id']
    preview=await teacher_client.get(base+'/slides/'+sid+'/preview')
    assert preview.status_code==200,preview.text
    assert 'data:image/png;base64,' in preview.text and '../images/' not in preview.text and 'private-original' not in preview.text
    export=await teacher_client.get(base+'/export');assert export.status_code==200,export.text
    with ZipFile(BytesIO(export.content)) as z:assert '<p:pic>' in z.read('ppt/slides/slide1.xml').decode()
    await make_user('reference-other');other=make_client(app)
    await other.post('/api/auth/login',json={'username':'reference-other','password':'pass1234'})
    assert (await other.get(base+'/slides/'+sid+'/preview')).status_code==404
    assert (await other.get(base+'/export')).status_code==404
    # A descriptor from another project must not authorize this file.
    alien=await other.post('/api/presentations',json={'title':'其他项目','prompt':'测试课件'})
    alien_id=UUID(alien.json()['id']);material.project_id=alien_id;await db.commit()
    assert (await teacher_client.get(base+'/slides/'+sid+'/preview')).status_code==409
    material.project_id=project.id;await db.commit()
    monkeypatch.setattr(master_assets,'parse_document',lambda *a,**k:[{'location':'图片1','text':'','image':png((200,200))}])
    assert (await teacher_client.get(base+'/export')).status_code==409
    saved=await teacher_client.put(base+'/outline',json={'revision':1,'slides':[slide]});assert saved.status_code==200
    assert (await teacher_client.get(base+'/export')).status_code==409
    assert (await teacher_client.get(base+'/slides/'+sid+'/preview')).status_code==409


async def test_svg_page_retries_invalid_geometry_without_text_fallback():
    import pytest
    from tests.test_ppt_master_runtime import sample_svg
    from unittest.mock import patch
    sid=str(uuid4());slide={'id':sid,'layout':'concept','title':'光合作用','content':'观察叶片','slots':{}}
    design={'id':sid,'visual_kind':'diagram','visual_brief':'叶片和太阳','composition':'左文右图'}
    invalid=sample_svg().replace('</svg>','<script>alert(1)</script></svg>')
    model=AsyncMock(side_effect=[{'svg':invalid},{'svg':sample_svg()}])
    with patch.object(generation,'chat_json',model):
        page=await generation.render_page(object(),slide,{'slides':[slide]},[],design,{})
    assert page['svg']==sample_svg() and page['renderer']=='ppt-master-6.7.0'
    assert model.call_args.args[1][-1]['role']=='user'
    with patch.object(generation,'chat_json',AsyncMock(return_value={'svg':invalid})):
        with pytest.raises(generation.GenerationError,match='图文设计未通过'):await generation.render_page(object(),slide,{'slides':[slide]},[],design,{})


async def test_worker_never_publishes_unapproved_result_and_preserves_previous_deck(teacher_client,db,settings,monkeypatch):
    from app.presentations import router
    from tests.test_ppt_rendering import page
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']));old=page()
    project.outline=checked_outline({'slides':[old]});project.pages=[old];project.rendered_revision=1;project.confirmed_revision=1;project.status='ready';await db.commit()
    rejected={**old,'review':{'status':'needs_attention','issues':[{'message':'实验结果不正确'}]}}
    monkeypatch.setattr(generation,'build_result',AsyncMock(return_value={'outline':project.outline,'pages':[rejected],'warnings':[]}))
    base='/api/presentations/'+p['id']
    started=await teacher_client.post(base+'/jobs/render',json={'revision':1,'idempotency_key':'unapproved-publish'})
    assert started.status_code==200,started.text
    await tasks.run_job(started.json()['id'],settings)
    state=(await teacher_client.get(base)).json()
    assert state['job']['status']=='failed'
    assert '文件验证未完成' in state['job']['error_message']
    assert state['revision']==state['rendered_revision']==1
    assert state['pages']==[old]


async def test_partial_export_contains_only_approved_pages_and_full_export_stays_blocked(teacher_client,db):
    from app.presentations.quality import VERSION,content_digest
    from tests.test_ppt_master_runtime import sample_svg
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    slides=[{'id':str(uuid4()),'layout':'concept','title':title,'content':'植物利用光制造有机物','slots':{}} for title in ['通过页','失败页']]
    good={**slides[0],'renderer':'ppt-master-6.7.0','svg':sample_svg(),'design':{'visual_brief':'叶片','composition':'图文并排'}}
    good['review']={'status':'passed','issues':[],'repairs':0,'version':VERSION,'content_sha256':content_digest(good),'image_sha256':'a'*64}
    bad={**slides[1],'review':{'status':'needs_attention','issues':[{'category':'visual','message':'重叠','repair':'重绘'}],'repairs':2}}
    project.outline=checked_outline({'slides':slides});project.pages=[good,bad];project.status='partial_ready';project.rendered_revision=1;await db.commit()
    base='/api/presentations/'+p['id']
    assert (await teacher_client.get(base+'/export')).status_code==409
    assert (await teacher_client.get(base+'/slides/'+good['id']+'/preview')).status_code==200
    assert (await teacher_client.get(base+'/slides/'+bad['id']+'/preview')).status_code==409
    exported=await teacher_client.get(base+'/export?partial=true')
    assert exported.status_code==200,exported.text
    assert len(Presentation(BytesIO(exported.content)).slides)==1
    from urllib.parse import unquote
    assert '部分完成' in unquote(exported.headers['content-disposition'])
    good=json.loads(json.dumps(good));good['svg']=good['svg'].replace('光合作用','更改内容');await db.execute(update(PptProject).where(PptProject.id==project.id).values(pages=[good,bad]));await db.commit()
    assert (await teacher_client.get(base+'/export?partial=true')).status_code==409


async def test_failed_first_page_does_not_abort_later_pages_and_retry_keeps_approved(monkeypatch):
    from types import SimpleNamespace
    from tests.test_ppt_master_runtime import sample_svg
    from app.presentations.quality import require_approved_pages
    slides=[{'id':str(uuid4()),'layout':'concept','title':title,'content':'植物利用光制造有机物','slots':{}} for title in ['失败页','通过页']]
    project=SimpleNamespace(outline=checked_outline({'slides':slides}),pages=[],title='独立恢复')
    job=SimpleNamespace(kind='render',payload={'slide_ids':[],'instruction':''},checkpoint={})
    reference={'name':'a'*24+'.png'}
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([{'image':reference}],[])))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    async def assets(db,p,settings,needed=None,*,pages=None):
        assert needed is not None,'load only each page required assets'
        return {reference['name']:b'authorized'} if reference['name'] in needed else {}
    monkeypatch.setattr(generation,'load_assets',assets)
    designs=[{'id':s['id'],'visual_kind':'diagram','visual_brief':'叶片','composition':'图文并排'} for s in slides]
    monkeypatch.setattr(generation,'plan_pages',AsyncMock(return_value={'pages':designs}))
    attempted=[]
    async def render_page(config,slide,*args,**kwargs):
        attempted.append(slide['id'])
        assert reference['name'] in args[-1],'authorized references must load before SVG validation'
        if slide['id']==slides[0]['id']:raise generation.GenerationError('不支持的属性：stroke-dasharray')
        return {**slide,'svg':sample_svg(),'renderer':'ppt-master-6.7.0','design':designs[1]}
    monkeypatch.setattr(generation,'render_page',render_page)
    monkeypatch.setattr(generation.master,'compile_deck',lambda *args:b'PK-test')
    monkeypatch.setattr(generation,'render_review_images',lambda data,count:[b'final-page']*count)
    monkeypatch.setattr(generation,'review_page',AsyncMock(return_value=[]))
    async def progress(value,stage,checkpoint=None):
        if checkpoint is not None:job.checkpoint=json.loads(json.dumps(checkpoint))
    options={'settings':SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA='),'progress':progress}
    result=await generation.build_result(None,project,job,**options)
    assert len(result['pages'])==2 and result['pages'][0].get('fallback')
    assert result['partial'] is False
    require_approved_pages(result['pages'])
    require_approved_pages([result['pages'][1]])
    assert slides[1]['id'] in attempted
    # Recovering a worker must reuse exact-approved checkpoint content.
    reviewed=generation.review_page.await_count
    resumed=await generation.build_result(None,project,job,**options)
    assert resumed['pages'][1]==result['pages'][1] and generation.review_page.await_count==reviewed
    project.pages=result['pages'];job.checkpoint={};attempted.clear()
    again=await generation.build_result(None,project,job,**options)
    assert slides[1]['id'] not in attempted and again['pages'][1]==project.pages[1]


async def test_failed_selected_revision_preserves_previous_complete_deck(teacher_client,db,settings,monkeypatch):
    from app.presentations import router
    from app.presentations.quality import VERSION,content_digest
    from tests.test_ppt_master_runtime import sample_svg
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    slides=[{'id':str(uuid4()),'layout':'concept','title':str(n),'content':'教学内容','slots':{}} for n in range(2)]
    pages=[]
    for slide in slides:
        page={**slide,'renderer':'ppt-master-6.7.0','svg':sample_svg(),'design':{'visual_brief':'叶片','composition':'图文并排'}}
        page['review']={'status':'passed','issues':[],'repairs':0,'version':VERSION,'content_sha256':content_digest(page),'image_sha256':'a'*64};pages.append(page)
    project.outline=checked_outline({'slides':slides});project.pages=pages;project.status='ready';project.rendered_revision=1;project.confirmed_revision=1;await db.commit()
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    rejected={**slides[0],'review':{'status':'needs_attention','repairs':2,'issues':[]}}
    monkeypatch.setattr(generation,'build_result',AsyncMock(return_value={'outline':project.outline,'pages':[rejected,pages[1]],'warnings':[],'partial':True}))
    base='/api/presentations/'+p['id']
    job=await teacher_client.post(base+'/jobs/revise',json={'revision':1,'idempotency_key':'preserve-full','slide_ids':[slides[0]['id']],'instruction':'修改本页'})
    assert job.status_code==200,job.text
    await tasks.run_job(job.json()['id'],settings)
    state=(await teacher_client.get(base)).json()
    assert state['job']['status']=='failed' and state['pages']==pages and state['revision']==1


async def test_render_requires_knowledge_check_of_exact_saved_outline(teacher_client,db,monkeypatch):
    from app.presentations import router
    from app.presentations.knowledge import outline_digest,VERSION
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    p=await create(teacher_client);base='/api/presentations/'+p['id']
    project=await db.get(PptProject,UUID(p['id']))
    slide={'id':str(uuid4()),'layout':'concept','title':'对称轴','content':'x=-b/(2a)，a≠0','slots':{}}
    project.outline={'slides':[slide]};await db.commit()
    r=await teacher_client.post(base+'/jobs/render',json={'revision':1,'idempotency_key':'unchecked'})
    assert r.status_code==409 and r.json()['detail']['code']=='outline_review_required'
    checked={'slides':[slide]}
    checked['review']={'version':VERSION,'status':'passed','issues':[],'changes':[],'content_sha256':outline_digest(checked)}
    project.outline=checked;await db.commit()
    saved=await teacher_client.put(base+'/outline',json={'revision':1,'slides':[slide]})
    assert saved.json()['outline']['review']['status']=='passed'
    edited={**slide,'content':'x=b/(2a)'}
    saved=await teacher_client.put(base+'/outline',json={'revision':2,'slides':[edited]})
    assert not saved.json()['outline'].get('review')
    r=await teacher_client.post(base+'/jobs/render',json={'revision':3,'idempotency_key':'changed'})
    assert r.status_code==409
    r=await teacher_client.post(base+'/jobs/check',json={'revision':3,'idempotency_key':'check'})
    assert r.status_code==200


async def test_unchanged_knowledge_check_preserves_existing_deck_and_partial_pages(teacher_client,db,settings,monkeypatch):
    from app.presentations import router
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    slide={'id':str(uuid4()),'layout':'concept','title':'定义','content':'a≠0','slots':{}}
    project.outline={'slides':[slide]};project.pages=[{**slide,'review':{'status':'needs_attention','issues':[]}}];project.rendered_revision=1;project.status='partial_ready';await db.commit()
    monkeypatch.setattr(generation,'build_result',AsyncMock(return_value={'outline':checked_outline(project.outline),'pages':[],'warnings':[]}))
    r=await teacher_client.post('/api/presentations/'+p['id']+'/jobs/check',json={'revision':1,'idempotency_key':'keep-deck'})
    await tasks.run_job(r.json()['id'],settings)
    state=(await teacher_client.get('/api/presentations/'+p['id'])).json()
    assert len(state['pages'])==1 and state['rendered_revision']==state['revision']

async def test_full_export_with_visual_suggestions_is_completed_and_downloadable(teacher_client,db):
    from app.presentations.quality import content_digest,VERSION
    from tests.test_ppt_master_runtime import formula_svg
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    slide={'id':str(uuid4()),'layout':'concept','title':'对称轴','content':'x=-b/(2a)，a≠0','slots':{}}
    page={**slide,'renderer':'ppt-master-6.7.0','svg':formula_svg(),'design':{'visual_kind':'text','visual_brief':'公式','composition':'纯文字'}}
    page['review']={'version':VERSION,'status':'passed','issues':[{'category':'visual','severity':'suggestion','message':'留白可增加','repair':'增加段距'}],'content_sha256':content_digest(page),'image_sha256':'a'*64}
    project.outline=checked_outline({'slides':[slide]});project.pages=[page];project.rendered_revision=1;project.status='ready';await db.commit()
    base='/api/presentations/'+p['id']
    assert (await teacher_client.get(base+'/slides/'+slide['id']+'/preview')).status_code==200
    exported=await teacher_client.get(base+'/export')
    assert exported.status_code==200 and len(Presentation(BytesIO(exported.content)).slides)==1
    assert p['id'] in [x['id'] for x in (await teacher_client.get('/api/presentations?scope=completed')).json()['items']]
    assert p['id'] not in [x['id'] for x in (await teacher_client.get('/api/presentations?scope=drafts')).json()['items']]


@pytest.mark.parametrize('selected_renderer',[None,'ppt-master-old'])
async def test_revise_rebuilds_selected_failed_page_without_touching_approved(teacher_client,db,settings,monkeypatch,selected_renderer):
    from app.presentations import router
    from app.presentations.quality import VERSION,content_digest
    from tests.test_ppt_master_runtime import sample_svg
    monkeypatch.setattr(router,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([],[])))
    monkeypatch.setattr(generation,'load_assets',AsyncMock(return_value={}))
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    slides=[{'id':str(uuid4()),'layout':'concept','title':title,'content':'植物利用光制造有机物','slots':{}} for title in ['已通过页','制作失败页']]
    good={**slides[0],'renderer':generation.master.VERSION,'svg':sample_svg(),'design':{'visual_kind':'diagram','visual_brief':'叶片','composition':'图文并排'}}
    good['review']={'status':'passed','issues':[],'repairs':0,'version':VERSION,'content_sha256':content_digest(good),'image_sha256':'a'*64}
    failed={**slides[1],'review':{'status':'needs_attention','repairs':0,'issues':[{'category':'generation','severity':'blocking','message':'页面制作失败','repair':'重建所选页'}]}}
    if selected_renderer:failed.update(renderer=selected_renderer,svg=sample_svg())
    project.outline=checked_outline({'slides':slides});project.pages=[good,failed];project.status='partial_ready';project.rendered_revision=project.confirmed_revision=1;await db.commit()
    async def model(config,messages,**kwargs):
        content=messages[-1]['content']
        request=json.loads(content[0]['text'] if isinstance(content,list) else content)
        if request['operation']=='revision_review':return {'status':'passed','issues':[]}
        if request['operation']=='edit_svg':
            from tests.test_ppt_master_runtime import sample_svg
            return {'svg':sample_svg().replace('光合作用','修改后的标题'),'slide':{**request['slide'],'title':'修改后的标题','content':'加入叶绿体实验'}}
        if request['operation']=='rewrite':return {'slides':[{**slides[1],'content':'观察光照下植物制造有机物'}]}
        if request['operation']=='knowledge_review':return {'status':'passed','issues':[]}
        raise AssertionError(request['operation'])
    monkeypatch.setattr(generation,'chat_json',model)
    async def plan(config,selected,*args):
        assert [s['id'] for s in selected]==[slides[1]['id']]
        return {'pages':[{'id':slides[1]['id'],'visual_kind':'text','visual_brief':'说明知识','composition':'清晰正文'}]}
    monkeypatch.setattr(generation,'plan_pages',plan)
    async def render_page(config,slide,*args,**kwargs):
        assert slide['id']==slides[1]['id'],'approved page must never be regenerated'
        return {**slide,'renderer':generation.master.VERSION,'svg':sample_svg(),'design':{'visual_kind':'text','visual_brief':'说明知识','composition':'清晰正文'}}
    monkeypatch.setattr(generation,'render_page',render_page)
    monkeypatch.setattr(generation.master,'compile_deck',lambda *args:b'PK-test')
    monkeypatch.setattr(generation,'render_review_images',lambda data,count:[b'rendered-page']*count)
    monkeypatch.setattr(generation,'review_page',AsyncMock(return_value=[]))
    base='/api/presentations/'+p['id']
    job=await teacher_client.post(base+'/jobs/revise',json={'revision':1,'idempotency_key':'revise-failed-placeholder','slide_ids':[slides[1]['id']],'instruction':'重做所选失败页，保留其他页面'})
    assert job.status_code==200,job.text
    await tasks.run_job(job.json()['id'],settings)
    state=(await teacher_client.get(base)).json()
    assert state['job']['status']=='succeeded',state['job']['error_message']
    assert state['status']=='ready' and state['revision']==state['rendered_revision']==2
    assert state['pages'][0]==good
    assert state['pages'][1]['review']['status']=='passed'
    assert state['pages'][1]['renderer']==generation.master.VERSION
    assert [page['id'] for page in state['pages']]==[s['id'] for s in slides]


async def test_revise_still_rejects_retaining_unselected_legacy_renderer(monkeypatch):
    from types import SimpleNamespace
    slides=[{'id':str(uuid4()),'layout':'concept','title':str(n),'content':'教学内容','slots':{}} for n in range(2)]
    project=SimpleNamespace(outline=checked_outline({'slides':slides}),pages=[{**s,'renderer':'old','svg':'legacy'} for s in slides],title='旧版')
    job=SimpleNamespace(kind='revise',payload={'slide_ids':[slides[1]['id']],'instruction':'重做'},checkpoint={})
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([],[])))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'load_assets',AsyncMock(return_value={}))
    with pytest.raises(generation.GenerationError,match='旧版课件'):
        await generation.build_result(None,project,job,SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA='),AsyncMock())
