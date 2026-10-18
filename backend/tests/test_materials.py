import io
import zipfile
import pytest
from app.materials.parsers import validate_document, parse_document

def test_zip_bomb_and_disguised_extension_rejected():
    with pytest.raises(ValueError):
        validate_document('bad.pdf', b'not a pdf')
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w',compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('../escape',b'test')
    with pytest.raises(ValueError):
        validate_document('bad.docx',stream.getvalue())

def test_docx_keeps_paragraph_locations():
    from docx import Document
    doc=Document();doc.add_paragraph('第一段教材');doc.add_paragraph('第二段教材')
    stream=io.BytesIO();doc.save(stream)
    chunks=parse_document('test.docx',stream.getvalue())
    assert chunks[0]['text']=='第一段教材'
    assert chunks[1]['location']=='第 2 段'

def test_docx_skips_embedded_images_with_unsupported_format(monkeypatch):
    import types
    from docx import Document
    stream=io.BytesIO(); Document().save(stream)
    fake_document=types.SimpleNamespace(
        paragraphs=[types.SimpleNamespace(text='正文内容')],
        tables=[],
        part=types.SimpleNamespace(rels={
            'image': types.SimpleNamespace(
                is_external=False,
                reltype='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image',
                target_part=types.SimpleNamespace(blob=b'unsupported-image-format'),
            ),
        }),
    )
    monkeypatch.setattr('docx.Document', lambda _: fake_document)
    warnings=[]
    chunks=parse_document('test.docx',stream.getvalue(),warnings=warnings)
    assert [chunk['text'] for chunk in chunks] == ['正文内容']
    assert all('image' not in chunk for chunk in chunks)
    assert len(warnings)==1
    assert '文档插图' in warnings[0]['location']
    assert '未解析' in warnings[0]['message']
    assert 'PNG/JPEG/WebP' in warnings[0]['message']


@pytest.mark.parametrize('extension',['docx','pptx'])
def test_unsupported_embedded_image_warns_but_preserves_text(extension):
    from PIL import Image
    picture=io.BytesIO();Image.new('RGB',(8,8)).save(picture,format='BMP');picture.seek(0)
    if extension=='docx':
        from docx import Document
        document=Document();document.add_paragraph('电路正文');document.add_picture(picture)
    else:
        from pptx import Presentation
        from pptx.util import Inches
        document=Presentation();slide=document.slides.add_slide(document.slide_layouts[1])
        slide.shapes.title.text='电路正文';slide.shapes.add_picture(picture,0,0,width=Inches(1))
    stream=io.BytesIO();document.save(stream)
    warnings=[]
    chunks=parse_document('电路.'+extension,stream.getvalue(),warnings=warnings)
    assert any('电路正文' in chunk['text'] for chunk in chunks)
    assert len(warnings)==1 and warnings[0]['location']
    assert all('image' not in chunk for chunk in chunks)


async def test_parse_warnings_persist_and_are_exposed_on_material_and_outline(db,teacher_client,settings,monkeypatch):
    from uuid import UUID
    from docx import Document
    from PIL import Image
    from app.models import VideoProject,Material
    from app.jobs.service import create_job
    from app.jobs.tasks import run_job
    data=(await teacher_client.post('/api/projects',json={'prompt':'电路'})).json()
    project=await db.get(VideoProject,UUID(data['id']))
    material=Material(owner_id=project.owner_id,project_id=project.id,filename='电路.docx',
        object_key='private/test',content_type='application/octet-stream',size=100)
    db.add(material);await db.flush()
    image=io.BytesIO();Image.new('RGB',(8,8)).save(image,format='BMP');image.seek(0)
    doc=Document();doc.add_paragraph('串联电路说明');doc.add_picture(image)
    stream=io.BytesIO();doc.save(stream)
    class Storage:
        def __init__(self,settings): pass
        def get(self,key): return stream.getvalue()
    monkeypatch.setattr('app.jobs.tasks.MaterialStorage',Storage)
    job=await create_job(db,project,'parse','warnings',{'material_id':str(material.id)})
    await db.commit();await run_job(job.id,settings)
    response=(await teacher_client.get(f'/api/projects/{project.id}/materials')).json()[0]
    assert response['status']=='ready'
    assert len(response['parse_warnings'])==1
    outline=(await teacher_client.get(f'/api/projects/{project.id}/outline')).json()
    assert outline['material_warnings'][0]['filename']=='电路.docx'
    assert '未解析' in outline['material_warnings'][0]['message']

async def test_upload_requires_same_origin_and_owned_project(teacher_client,admin_client):
    project=(await teacher_client.post('/api/projects',json={'prompt':'数学'})).json()
    path='/api/projects/'+project['id']+'/materials'
    response=await teacher_client.post(path,files={'file':('a.pdf',b'%PDF-invalid','application/pdf')})
    assert response.status_code==403
    response=await teacher_client.post(path,headers={'Origin':'http://test:5173'},files={'file':('a.pdf',b'%PDF-invalid','application/pdf')})
    assert response.status_code==403
    response=await admin_client.post(path,headers={'Origin':'http://test'},files={'file':('a.pdf',b'%PDF-invalid','application/pdf')})
    assert response.status_code==404

async def test_deferred_upload_is_persisted_without_parse_job(teacher_client,monkeypatch):
    import io
    from docx import Document
    doc=Document();doc.add_paragraph('待生成大纲的资料');stream=io.BytesIO();doc.save(stream)
    objects={}
    class Storage:
        def __init__(self,settings):pass
        def put(self,key,data,mime):objects[key]=data
    monkeypatch.setattr('app.materials.router.MaterialStorage',Storage)
    p=(await teacher_client.post('/api/projects',json={'prompt':'数学'})).json()
    path='/api/projects/'+p['id']+'/materials'
    response=await teacher_client.post(path+'?defer_parse=true',headers={'Origin':'http://test'},files={'file':('课件.docx',stream.getvalue())})
    assert response.status_code==201,response.text
    assert response.json()['status']=='stored' and response.json()['job_id'] is None
    assert len(objects)==1
    prepared=await teacher_client.post(path+'/prepare')
    assert prepared.status_code==202,prepared.text
    assert len(prepared.json()['job_ids'])==1

def test_pdf_scan_has_page_image_and_encrypted_rejected():
    import pymupdf
    doc=pymupdf.open();doc.new_page()
    data=doc.tobytes()
    chunks=parse_document('扫描.pdf',data)
    assert chunks[0]['location']=='第 1 页（图像）' and chunks[0]['image']
    encrypted=doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='reader')
    with pytest.raises(ValueError,match='密码'):validate_document('加密.pdf',encrypted)

def test_pptx_has_slide_location():
    from pptx import Presentation
    presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[1]);slide.shapes.title.text='二次函数'
    stream=io.BytesIO();presentation.save(stream)
    chunks=parse_document('课件.pptx',stream.getvalue())
    assert chunks[0]['location']=='第 1 张幻灯片' and '二次函数' in chunks[0]['text']

def test_legacy_converter_timeout_terminates_process_group(monkeypatch,tmp_path):
    import subprocess
    from types import SimpleNamespace
    from unittest.mock import Mock
    from app.materials import parsers
    process=Mock(pid=1234)
    process.wait.side_effect=[subprocess.TimeoutExpired('libreoffice',70),0]
    kill=Mock()
    monkeypatch.setattr(parsers,'os',SimpleNamespace(name='posix',killpg=kill))
    monkeypatch.setattr(parsers,'signal',SimpleNamespace(SIGKILL=9))
    monkeypatch.setattr(parsers.subprocess,'Popen',Mock(return_value=process))
    with pytest.raises(ValueError,match='转换超时'):
        parsers.convert_legacy('test.doc',b'synthetic',str(tmp_path))
    kill.assert_called_once_with(1234,parsers.signal.SIGKILL)
    assert process.wait.call_count==2
