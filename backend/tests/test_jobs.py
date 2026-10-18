import pytest
from fastapi import HTTPException

from app.jobs.service import count_active_generation_jobs, create_job, require_generation_capacity, update_job_progress
from app.models import VideoProject
from app.settings_store import GenerationParams, put_generation_params
from uuid import UUID

async def test_job_idempotency_and_owner_access(db,teacher_client,admin_client):
    project=(await teacher_client.post('/api/projects',json={'prompt':'数学'})).json()
    p=await db.get(VideoProject,UUID(project['id']))
    first=await create_job(db,p,'parse','token-1',{})
    second=await create_job(db,p,'parse','token-1',{})
    assert first.id==second.id
    await db.commit()
    assert (await teacher_client.get('/api/jobs/'+str(first.id))).status_code==200
    assert (await admin_client.get('/api/jobs/'+str(first.id))).status_code==404


async def test_job_progress_is_persisted_for_reopened_outline(db, teacher_client):
    from datetime import UTC, datetime, timedelta
    from uuid import uuid4

    project_data = (await teacher_client.post('/api/projects', json={'prompt':'数学'})).json()
    project = await db.get(VideoProject, UUID(project_data['id']))
    job = await create_job(db, project, 'generate', 'progress', {})
    job.status = 'running'
    job.lease_token = uuid4()
    job.lease_until = datetime.now(UTC) + timedelta(minutes=1)
    await db.commit()

    await update_job_progress(db, job.id, job.lease_token, 42, '审核大纲')
    await db.refresh(job)

    assert job.progress == 42
    assert job.stage == '审核大纲'


async def test_progress_does_not_commit_unpublished_project_output(db, teacher_client):
    from uuid import uuid4
    data = (await teacher_client.post('/api/projects', json={'prompt': '数学'})).json()
    project = await db.get(VideoProject, UUID(data['id']))
    job = await create_job(db, project, 'generate', 'isolated-progress', {})
    job.status, job.lease_token = 'running', uuid4()
    await db.commit()
    original = project.title
    project.title = '不能提前发布的内容'
    await update_job_progress(db, job.id, job.lease_token, 45, '生成大纲')
    await db.rollback()
    await db.refresh(project)
    assert project.title == original


async def test_progress_is_monotonic_and_rejects_stale_lease(db, teacher_client):
    from uuid import uuid4
    data = (await teacher_client.post('/api/projects', json={'prompt': '数学'})).json()
    project = await db.get(VideoProject, UUID(data['id']))
    job = await create_job(db, project, 'generate', 'monotonic-progress', {})
    token = uuid4()
    job.status, job.lease_token, job.progress = 'running', token, 80
    await db.commit()
    await update_job_progress(db, job.id, token, 25, '恢复任务')
    await db.refresh(job)
    assert job.progress == 80
    await db.commit()
    with pytest.raises(ValueError, match='接管|终止'):
        await update_job_progress(db, job.id, uuid4(), 90, '旧任务')
    await db.refresh(job)
    assert job.stage == '恢复任务'


async def test_generation_capacity_is_independent_per_teacher_and_excludes_parse_jobs(db, make_user):
    await put_generation_params(db, GenerationParams(max_concurrent_jobs=2))
    teacher_one = await make_user('capacity-teacher-1')
    teacher_two = await make_user('capacity-teacher-2')
    teacher_one_id, teacher_two_id = teacher_one.id, teacher_two.id
    first = VideoProject(owner_id=teacher_one.id, title='一', prompt='一')
    second = VideoProject(owner_id=teacher_one.id, title='二', prompt='二')
    other = VideoProject(owner_id=teacher_two.id, title='三', prompt='三')
    db.add_all([first, second, other])
    await db.flush()

    first_job = await create_job(db, first, 'generate', 'one', {})
    second_job = await create_job(db, second, 'generate', 'two', {})
    second_job.status = 'running'
    await create_job(db, first, 'parse', 'material', {})
    await create_job(db, other, 'generate', 'other', {})
    await db.commit()

    assert await count_active_generation_jobs(db, teacher_one_id) == 2
    assert await count_active_generation_jobs(db, teacher_one_id, exclude_job_id=first_job.id) == 1
    assert await count_active_generation_jobs(db, teacher_two_id) == 1

    with pytest.raises(HTTPException) as exc_info:
        await require_generation_capacity(db, teacher_one_id, 2)
    assert exc_info.value.status_code == 409
    assert exc_info.value.detail == {
        'code': 'generation_capacity_exceeded',
        'message': '超出最大任务数：已存在2个生成任务',
    }
    await db.rollback()

    # The second teacher still has its own available slot despite teacher one being full.
    assert await require_generation_capacity(db, teacher_two_id, 2) == 1
    await db.rollback()

async def test_parse_worker_is_durable_idempotent_and_recovers_expired_lease(db,teacher_client,settings,monkeypatch):
    import io
    from datetime import UTC,datetime,timedelta
    from docx import Document
    from sqlalchemy import select,func
    from app.models import Material,MaterialChunk,OutlineJob
    from app.jobs.tasks import run_job
    p=(await teacher_client.post('/api/projects',json={'prompt':'数学'})).json()
    project=await db.get(VideoProject,UUID(p['id']))
    material=Material(owner_id=project.owner_id,project_id=project.id,filename='课件.docx',object_key='private/key',content_type='application/octet-stream',size=100)
    db.add(material);await db.flush()
    doc=Document();doc.add_paragraph('第一段');doc.add_paragraph('第二段');stream=io.BytesIO();doc.save(stream)
    class Storage:
        def __init__(self,settings):pass
        def get(self,key):return stream.getvalue()
    monkeypatch.setattr('app.jobs.tasks.MaterialStorage',Storage)
    job=await create_job(db,project,'parse','recover',{'material_id':str(material.id)})
    job.status='running';job.lease_token=__import__('uuid').uuid4();job.lease_until=datetime.now(UTC)-timedelta(seconds=5)
    await db.commit()
    await run_job(job.id,settings)
    await db.refresh(job);await db.refresh(material)
    assert job.status=='succeeded' and material.status=='ready'
    assert await db.scalar(select(func.count()).select_from(MaterialChunk).where(MaterialChunk.material_id==material.id))==2
    # Crashes exhausting the budget must also make the material retryable.
    job.status='running';job.attempts=99;job.lease_until=datetime.now(UTC)-timedelta(seconds=5);material.status='parsing'
    await db.commit();await run_job(job.id,settings);await db.refresh(job);await db.refresh(material)
    assert job.status=='failed' and material.status=='failed' and job.lease_until is None
    assert job.lease_token is None
    await db.commit();await run_job(job.id,settings)
    assert await db.scalar(select(func.count()).select_from(MaterialChunk).where(MaterialChunk.material_id==material.id))==2

async def test_revoked_worker_cannot_publish_material(db,teacher_client,settings,monkeypatch):
    from uuid import uuid4
    from sqlalchemy import update
    import pytest
    from app.models import Material,OutlineJob
    from app.jobs.tasks import parse_job
    from app.db import create_sessionmaker
    p=(await teacher_client.post('/api/projects',json={'prompt':'数学'})).json()
    project=await db.get(VideoProject,UUID(p['id']))
    material=Material(owner_id=project.owner_id,project_id=project.id,filename='a.docx',object_key='key',content_type='application/octet-stream',size=1)
    db.add(material);await db.flush()
    job=await create_job(db,project,'parse','revoked',{'material_id':str(material.id)})
    job.status='running';job.lease_token=uuid4();await db.commit()
    maker=create_sessionmaker(db.bind)
    async def revoke(*args, **kwargs):
        async with maker() as other:
            await other.execute(update(OutlineJob).where(OutlineJob.id==job.id).values(status='failed'))
            await other.commit()
        return [{'text':'教材','location':'第1段'}]
    monkeypatch.setattr('app.jobs.tasks.asyncio.to_thread',revoke)
    storage=type('Storage',(),{'get':lambda self,key:b'doc'})()
    with pytest.raises(ValueError,match='接管|终止'):
        await parse_job(db,job,settings,storage)


async def test_stale_parser_cannot_reset_completed_material_before_parsing(db,teacher_client,settings):
    from uuid import uuid4
    from sqlalchemy import update
    from app.models import Material,OutlineJob
    from app.jobs.tasks import parse_job
    from app.db import create_sessionmaker
    data=(await teacher_client.post('/api/projects',json={'prompt':'电路'})).json()
    project=await db.get(VideoProject,UUID(data['id']))
    material=Material(owner_id=project.owner_id,project_id=project.id,filename='a.docx',
        object_key='key',content_type='application/octet-stream',size=1,status='pending')
    db.add(material);await db.flush()
    job=await create_job(db,project,'parse','stale-before-start',{'material_id':str(material.id)})
    job.status,job.lease_token='running',uuid4()
    await db.commit()
    maker=create_sessionmaker(db.bind)
    async with maker() as other:
        await other.execute(update(OutlineJob).where(OutlineJob.id==job.id).values(status='succeeded',lease_token=uuid4()))
        await other.execute(update(Material).where(Material.id==material.id).values(status='ready'))
        await other.commit()
    storage=type('Storage',(),{'get':lambda self,key:b'not-a-document'})()
    with pytest.raises(ValueError):
        await parse_job(db,job,settings,storage)
    await db.rollback()
    await db.refresh(material)
    assert material.status=='ready'


async def test_success_timestamp_marks_completion_after_model_returns(db, teacher_client, settings, monkeypatch):
    from datetime import UTC, datetime
    from sqlalchemy import text
    from app.jobs.tasks import run_job
    data=(await teacher_client.post('/api/projects',json={'prompt':'数学'})).json()
    project=await db.get(VideoProject,UUID(data['id']))
    job=await create_job(db,project,'generate','completion-order',{})
    await db.commit()
    model_returned=[]
    async def execute(session,job,settings):
        await session.execute(text('SELECT 1'))
        model_returned.append(datetime.now(UTC))
    monkeypatch.setattr('app.outline.service.execute_outline_job',execute)
    await run_job(job.id,settings)
    response=await teacher_client.get('/api/jobs/'+str(job.id))
    assert response.status_code==200
    assert response.json()['status']=='succeeded'
    assert datetime.fromisoformat(response.json()['updated_at']) >= model_returned[0]
