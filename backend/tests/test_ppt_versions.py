from copy import deepcopy
from io import BytesIO
from uuid import UUID
from pptx import Presentation
from app.presentations.models import PptProject
from app.presentations.service import invalidate
from tests.test_ppt_api import create
from tests.test_ppt_rendering import page

async def test_previous_deck_keeps_own_pages_title_and_owner_checked_preview_export(teacher_client,db,make_user,make_client,app):
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    old=page();old['title']='旧版页面';old['slots'][next(iter(old['slots']))]='旧版页面';project.pages=[old];project.outline={'slides':[old]};project.rendered_revision=1
    old_title=project.title;invalidate(project)
    new={**old,'title':'新版页面','slots':{k:'新版页面' if v=='旧版页面' else v for k,v in old['slots'].items()}};project.pages=[new];project.outline={'slides':[new]};project.title='新版课件';project.revision=2;project.rendered_revision=2;await db.commit()
    base='/api/presentations/'+p['id']
    state=(await teacher_client.get(base)).json()
    assert state['versions'][0]['revision']==1
    assert state['versions'][0]['title']==old_title
    snapshot=await teacher_client.get(base+'/versions/1');assert snapshot.status_code==200,snapshot.text
    assert snapshot.json()['pages'][0]['title']=='旧版页面'
    old_preview=await teacher_client.get(base+'/slides/'+old['id']+'/preview?version=1')
    assert old_preview.status_code==200 and '旧版页面' in old_preview.text and '新版页面' not in old_preview.text
    current=await teacher_client.get(base+'/slides/'+old['id']+'/preview');assert '新版页面' in current.text
    export=await teacher_client.get(base+'/export?version=1');assert export.status_code==200
    assert '旧版页面' in ''.join(getattr(s,'text','') for s in Presentation(BytesIO(export.content)).slides[0].shapes)
    assert (await teacher_client.get(base+'/versions/99')).status_code==404
    assert (await teacher_client.get(base+'/slides/'+old['id']+'/preview?version=99')).status_code==404
    await make_user('version-other');other=make_client(app);await other.post('/api/auth/login',json={'username':'version-other','password':'pass1234'})
    for path in ['/versions/1','/export?version=1','/slides/'+old['id']+'/preview?version=1']:
        assert (await other.get(base+path)).status_code==404

def test_archive_deduplicates_revision_and_preserves_legacy_snapshots():
    from types import SimpleNamespace
    a=page();p=SimpleNamespace(pages=[a],outline={'slides':[a]},title='标题',warnings=[],history=[{'revision':1,'pages':[a]}],rendered_revision=1,confirmed_revision=1,updated_at=None)
    invalidate(p)
    assert len(p.history)==1 and p.history[0]['title']=='标题'

async def test_legacy_page_only_history_and_current_revision_are_resolved_exactly(teacher_client,db):
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    a=page();project.pages=[a];project.outline={'slides':[a]};project.revision=3;project.rendered_revision=3
    project.history=[{'revision':1,'pages':[a]}];await db.commit();base='/api/presentations/'+p['id']
    assert (await teacher_client.get(base+'/versions/1')).json()['revision']==1
    assert (await teacher_client.get(base+'/slides/'+a['id']+'/preview?version=1')).status_code==200
    assert (await teacher_client.get(base+'/versions/3')).json()['revision']==3
    project.rendered_revision=2;await db.commit()
    assert (await teacher_client.get(base+'/versions/3')).status_code==404

async def test_history_keeps_uploaded_image_after_material_removed_from_current_deck(teacher_client,db,monkeypatch):
    from app.presentations.models import PptMaterial
    from app.presentations.master_assets import normalize_image,image_descriptor
    from app.materials.storage import MaterialStorage
    from tests.test_ppt_master_assets import png
    from tests.test_ppt_master_runtime import sample_svg
    from uuid import uuid4
    from sqlalchemy import select
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    data=normalize_image(png());mid=uuid4();descriptor=image_descriptor(mid,0,data)
    material=PptMaterial(id=mid,project_id=project.id,owner_id=project.owner_id,filename='leaf.png',object_key='history-leaf',content_type='image/png',size=len(data),status='ready',chunks=[{'text':'叶片','parse_version':2,'image':descriptor}])
    db.add(material)
    a={**page(),'slots':{},'renderer':'ppt-master-6.7.0','svg':sample_svg().replace('</svg>',f'<image href="../images/{descriptor["name"]}" x="64" y="400" width="100" height="100"/></svg>')}
    a['design']={'id':a['id'],'visual_kind':'diagram','visual_brief':'叶片配图','composition':'保留上传图片'}
    project.pages=[a];project.outline={'slides':[{k:a[k] for k in ('id','layout','title','content','slots')}]};project.rendered_revision=project.revision;await db.commit()
    monkeypatch.setattr(MaterialStorage,'get',lambda *args:data)
    base='/api/presentations/'+p['id']
    removed=await teacher_client.delete(base+'/materials/'+str(mid)+'?revision=1');assert removed.status_code==204,removed.text
    assert (await teacher_client.get(base+'/materials')).json()==[]
    preview=await teacher_client.get(base+'/slides/'+a['id']+'/preview?version=1')
    assert preview.status_code==200,preview.text
    assert 'data:image/png;base64' in preview.text
    from app.presentations.master_runtime import compile_deck
    native=compile_deck([a],project.title,{descriptor['name']:data})
    assert len(Presentation(BytesIO(native)).slides)==1
    exported=await teacher_client.get(base+'/export?version=1');assert exported.status_code==200,exported.text
    assert len(Presentation(BytesIO(exported.content)).slides)==1
    assert (await teacher_client.get(base+'/materials/'+str(mid)+'/download')).status_code==404
    assert (await teacher_client.post(base+'/materials/'+str(mid)+'/retry?revision=2')).status_code==404
    from app.models import ObjectCleanup
    assert not (await db.scalars(select(ObjectCleanup).where(ObjectCleanup.object_key=='history-leaf'))).all()

async def test_archived_material_cleanup_waits_for_last_retained_reference(teacher_client,db):
    from app.presentations.models import PptMaterial
    from app.presentations.versions import prune_archived_materials
    from app.models import ObjectCleanup
    from sqlalchemy import select
    from uuid import uuid4
    p=await create(teacher_client);project=await db.get(PptProject,UUID(p['id']))
    name='abcdef0123456789abcdef01.png';mid=uuid4()
    material=PptMaterial(id=mid,project_id=project.id,owner_id=project.owner_id,filename='leaf.png',object_key='old-ref',content_type='image/png',size=100,status='archived',chunks=[{'image':{'name':name}}]);db.add(material);await db.flush()
    project.history=[];await prune_archived_materials(db,project);await db.flush()
    assert await db.get(PptMaterial,mid) is None
    assert await db.scalar(select(ObjectCleanup).where(ObjectCleanup.object_key=='old-ref'))

def test_retention_limits_old_versions_to_ten():
    from types import SimpleNamespace
    p=SimpleNamespace(pages=[],outline={},title='课件',warnings=[],history=[],rendered_revision=None,confirmed_revision=None,updated_at=None)
    for revision in range(1,14):
        p.pages=[page()];p.rendered_revision=revision;invalidate(p)
    assert [h['revision'] for h in p.history]==list(range(4,14))
