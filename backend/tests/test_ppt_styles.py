from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import AsyncMock
import json
from uuid import UUID
import pytest
from app.presentations import fallback,generation,master_runtime as master
from app.presentations.models import PptProject
from tests.test_ppt_fallback import outline

@pytest.mark.parametrize('style,name',[('lively-illustration','活泼插画 · 橙光'),('academic-formal','学术正式 · 墨红')])
async def test_approved_style_can_be_created_and_archived_without_relabeling_history(teacher_client,db,style,name):
    config=(await teacher_client.get('/api/presentations/config')).json()
    assert {t['id'] for t in config['templates']}=={'clear-blue-teaching','lively-illustration','academic-formal'}
    response=await teacher_client.post('/api/presentations',json={'title':'二次函数','prompt':'初三数学','style':style})
    assert response.status_code==200,response.text
    data=response.json();assert data['template_id']==style and data['template_name']==name
    project=await db.get(PptProject,UUID(data['id']));source=outline('读懂图像')
    project.pages=fallback.make_pages(source['slides'][0],style=style);project.outline=source;project.rendered_revision=1;await db.commit()
    saved=await teacher_client.put('/api/presentations/'+data['id'],json={'title':'新版','prompt':'新要求','style':'clear-blue-teaching','revision':1})
    assert saved.status_code==200,saved.text
    old=(await teacher_client.get('/api/presentations/'+data['id']+'/versions/1')).json()
    assert old['template_id']==style and old['template_name']==name
    assert old['pages'][0]['design']['template']['id']==style

@pytest.mark.parametrize('style,bg,title',[('lively-illustration','#FFF9ED','#243C36'),('academic-formal','#FAFAF7','#202D3A')])
def test_fallback_retains_selected_style_and_still_rejects_tampering(style,bg,title):
    source=outline('函数知识。'*180);sid=source['slides'][0]['id']
    pages=fallback.make_pages(source['slides'][0],style=style)
    assert bg in pages[0]['svg'] and title in pages[0]['svg']
    assert pages[0]['design']['template']=={'id':style,'version':'1.0.0'}
    result=fallback.paginated_outline(source,{sid:pages},style=style)
    assert ''.join(s['content'] for s in result['slides'])==source['slides'][0]['content']
    tampered=deepcopy(pages);tampered[0]['svg']=tampered[0]['svg'].replace(bg,'#FFFFFF')
    with pytest.raises(ValueError):fallback.paginated_outline(source,{sid:tampered},style=style)
    with pytest.raises(ValueError):fallback.paginated_outline(source,{sid:pages},style='clear-blue-teaching')

@pytest.mark.parametrize('style,accent',[('lively-illustration','#B84A26'),('academic-formal','#8B3547')])
async def test_planner_and_executor_use_approved_template_not_clearblue(monkeypatch,style,accent):
    from tests.test_ppt_master_runtime import sample_svg
    source=outline('使用图像解释二次函数');slide=source['slides'][0];requests=[]
    async def model(config,payload,**kwargs):
        requests.append(payload)
        if payload['operation']=='design':return {'pages':[{'id':slide['id'],'visual_kind':'diagram','visual_brief':'数学图像','composition':'图文并排'}]}
        return {'svg':sample_svg()}
    monkeypatch.setattr(generation,'call',model)
    monkeypatch.setattr(master,'check_page',lambda *a:None)
    plan=await generation.plan_pages(None,[slide],source,[],style=style)
    await generation.render_page(None,slide,source,[],plan['pages'][0],{})
    assert len(requests)==2
    for request in requests:
        assert request['template']['id']==style
        assert accent in json.dumps(request['template'])
    assert '{{FORMULA}}' in requests[1]['template']['layout_svg']
    assert 'data-pptx-layout' in requests[1]['template']['layout_svg']

@pytest.mark.parametrize('style,name,bg',[('lively-illustration','活泼插画 · 橙光','#FFF9ED'),('academic-formal','学术正式 · 墨红','#FAFAF7')])
def test_native_export_uses_selected_template_design_artifacts(style,name,bg):
    from io import BytesIO
    from pptx import Presentation
    pages=fallback.make_pages(outline('函数图像的开口与顶点。')['slides'][0],style=style)
    spec,lock=master.design_artifacts(pages,'教学样例',{})
    assert name in spec and bg in lock
    data=master.compile_deck(pages,'教学样例',{})
    deck=Presentation(BytesIO(data))
    assert len(deck.slides)==1
    import zipfile
    from xml.etree import ElementTree as ET
    with zipfile.ZipFile(BytesIO(data)) as archive:
        xml=ET.fromstring(archive.read('ppt/slides/slide1.xml'))
        assert '函数图像' in ''.join(xml.itertext())
    assert not any(shape.shape_type==13 for shape in deck.slides[0].shapes)

def test_auto_and_legacy_defaults_stay_clearblue():
    from app.presentations.schemas import ProjectInput
    assert ProjectInput(title='标题',prompt='需求').style=='auto'
    page=fallback.make_pages(outline('旧版内容')['slides'][0])[0]
    assert 'template' not in page['design']
    assert '#FFFFFF' in page['svg'] and '#173359' in page['svg']

@pytest.mark.parametrize('style,background',[('lively-illustration','#FFF9ED'),('academic-formal','#FAFAF7')])
async def test_render_resumes_only_matching_template_checkpoint(monkeypatch,style,background):
    source=outline('课堂知识')
    project=SimpleNamespace(outline=source,pages=[],title='模板恢复',style=style)
    wrong=fallback.make_pages(source['slides'][0])[0]
    job=SimpleNamespace(kind='render',payload={'slide_ids':[],'instruction':''},checkpoint={'renderer':master.VERSION,'outline':source,'pages':{wrong['id']:wrong},'planning_failed':'old failure'})
    monkeypatch.setattr(generation,'references',AsyncMock(return_value=([],[])))
    monkeypatch.setattr(generation,'resolve',AsyncMock(return_value=object()))
    monkeypatch.setattr(generation,'load_assets',AsyncMock(return_value={}))
    monkeypatch.setattr(generation,'plan_pages',AsyncMock(side_effect=generation.GenerationError('模型不可用')))
    monkeypatch.setattr(master,'compile_deck',lambda *a:b'PK-test')
    monkeypatch.setattr(generation,'render_review_images',lambda data,count:[b'image']*count)
    monkeypatch.setattr(generation,'review_page',AsyncMock(return_value=[]))
    async def progress(value,stage,checkpoint=None):
        if checkpoint is not None:job.checkpoint=deepcopy(checkpoint)
    result=await generation.build_result(None,project,job,SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA='),progress)
    assert result['pages'][0]['design']['template']['id']==style
    assert background in result['pages'][0]['svg']
    assert result['pages'][0].get('fallback')
    resumed=await generation.build_result(None,project,job,SimpleNamespace(master_key='AAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA='),progress)
    assert resumed==result

async def test_new_style_infographic_uses_native_template_and_direct_edits_keep_identity(monkeypatch):
    from tests.test_ppt_master_runtime import sample_svg
    source=outline('比较两种函数');slide=source['slides'][0]
    design={'id':slide['id'],'visual_kind':'infographic','visual_brief':'比较','composition':'左右并排','template':{'id':'academic-formal','version':'1.0.0'},'graphic':{'kind':'compare','items':[{'label':'甲','desc':'上升','motif':'none'},{'label':'乙','desc':'下降','motif':'none'}]}}
    async def model(config,payload,**kwargs):
        assert payload['template']['id']=='academic-formal'
        assert '{{LEFT_LABEL}}' in payload['template']['layout_svg']
        if payload['operation']=='edit_svg':return {'svg':sample_svg(),'slide':slide}
        return {'svg':sample_svg()}
    slide['layout']='compare'
    monkeypatch.setattr(generation,'call',model)
    monkeypatch.setattr(master,'check_page',lambda *a:None)
    page=await generation.render_page(None,slide,source,[],design,{})
    assert page['design']['visual_kind']=='diagram' and not page.get('graphics')
    edited=await generation.render_page(None,slide,source,[],design,{},current=page,instruction='只改标题位置')
    assert edited['design']['template']=={'id':'academic-formal','version':'1.0.0'}

@pytest.mark.parametrize('style',['lively-illustration','academic-formal'])
def test_multi_page_template_typography_compiles(style):
    from io import BytesIO
    from pptx import Presentation
    from uuid import uuid4
    pages=[]
    for index in range(3):
        slide={**outline('完整保留知识与公式')['slides'][0],'id':str(uuid4())}
        page=fallback.make_pages(slide,style=style)[0]
        page['svg']=page['svg'].replace('font-size="22"','font-size="18"')
        pages.append(page)
    spec,lock=master.design_artifacts(pages,'多页模板字号',{})
    data=master.compile_deck(pages,'多页模板字号',{})
    assert len(Presentation(BytesIO(data)).slides)==3
