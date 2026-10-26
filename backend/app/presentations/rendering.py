"""Render the approved template without executing model-authored markup or code.
SVG preview and editable PPTX use the same resolved geometry and wrapped lines.
"""
from functools import lru_cache
from io import BytesIO
import json
import math
from pathlib import Path
import unicodedata
import xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from .schemas import Slide

ROOT=Path(__file__).with_name('templates')
NS='http://www.w3.org/2000/svg'
ET.register_namespace('',NS)

def tag(node):return node.tag.rsplit('}',1)[-1]
def element(name,**attrs):return ET.Element('{'+NS+'}'+name,{k.replace('_','-'):str(v) for k,v in attrs.items()})

# CJK fonts render these East Asian ambiguous characters at full width.
# Counting them as Latin punctuation can make native PPTX text overrun its slot.
FULL_WIDTH_PUNCTUATION = '“”‘’—…'
CLOSING_PUNCTUATION = '。，、；：！？）］｝》〉」』】〕〗〙〛”’'
OPENING_PUNCTUATION = '（［｛《〈「『【〔〖〘〚“‘'

def units(value):
    return sum(1 if unicodedata.east_asian_width(c) in 'WF'
               or c in FULL_WIDTH_PUNCTUATION else .6 for c in value)

def wrap(value, width, size):
    lines = []
    for raw in value.splitlines() or ['']:
        if not raw:
            lines.append('')
            continue
        start = 0
        while start < len(raw):
            end = start
            while end < len(raw) and units(raw[start:end + 1]) * size <= width:
                end += 1
            end = max(start + 1, end)
            # Move the boundary back instead of overflowing a slot: closing
            # punctuation stays with preceding text, opening with following.
            while end > start + 1 and end < len(raw) and (
                raw[end] in CLOSING_PUNCTUATION
                or raw[end - 1] in OPENING_PUNCTUATION
            ):
                end -= 1
            lines.append(raw[start:end])
            start = end
    return lines

@lru_cache
def catalog():
    tokens=json.loads((ROOT/'tokens.json').read_text(encoding='utf-8-sig'))
    result={}
    for item in tokens['layouts']:
        source=ROOT/Path(item['source']).name
        tree=ET.parse(source).getroot();slots=[]
        for g in tree:
            if not g.get('data-pptx-placeholder') or g.get('id') in ['course-footer','slide-number']:continue
            bounds=[float(v) for v in g.get('data-pptx-bounds').split()]
            carrier=None if g.get('data-pptx-placeholder')=='object' else next((n for n in g if tag(n)=='text'),None)
            size=float(carrier.get('font-size','28')) if carrier is not None else 28
            max_lines=max(1,math.floor((bounds[3]-size*1.05)/(size*1.45))+1)
            slots.append({'id':g.get('id'),'width':bounds[2],'height':bounds[3],'font_size':size,'max_lines':max_lines,'max_units_per_line':round(bounds[2]/size,1),'visual':carrier is None})
        result[item['id']]={'name':item['name'],'source':source.name,'slots':slots}
    return result

def validate_page(page):
    page=Slide.model_validate(page).model_dump(mode='json')
    spec=catalog()[page['layout']]
    if set(page['slots'])!={s['id'] for s in spec['slots']}:raise ValueError('页面内容字段不完整，请重新生成此页')
    for slot in spec['slots']:
        value=page['slots'][slot['id']]
        if not value.strip() or len(value)>2000 or any(ord(c)<32 and c not in '\n\t' for c in value):raise ValueError('页面包含空内容或非法字符')
        if len(wrap(value,slot['width'],slot['font_size']))>slot['max_lines']:
            raise ValueError('页面内容超出版式容量，请精简或拆页：'+slot['id'])
    return page

def resolved_svg(page,course,number):
    page=validate_page(page)
    root=ET.parse(ROOT/catalog()[page['layout']]['source']).getroot()
    for group in list(root):
        if not group.get('data-pptx-placeholder'):continue
        ident=group.get('id');x,y,w,h=[float(v) for v in group.get('data-pptx-bounds').split()]
        value=course[:50] if ident=='course-footer' else str(number) if ident=='slide-number' else page['slots'][ident]
        carrier=None if group.get('data-pptx-placeholder')=='object' else next((n for n in group if tag(n)=='text'),None)
        if carrier is None:
            group.clear()
            group.append(element('rect',x=x,y=y,width=w,height=h,rx=8,fill='#EAF1FC'))
            carrier=element('text',x=x+w/2,y=y+h/2,font_size=28,fill='#2465D9',text_anchor='middle')
            group.append(carrier);size=28
        else:
            size=float(carrier.get('font-size','28'))
            for child in list(carrier):carrier.remove(child)
        lines=wrap(value,w,size)
        # Resolve each line as an independent text run for identical SVG/PPTX positions.
        start=(y+h/2-size*1.45*(len(lines)-1)/2+size*.35) if carrier.get('text-anchor')=='middle' and ident not in ['slide-number'] else y+min(h-size*1.45*(len(lines)-1),size*1.12)
        cx=float(carrier.get('x',x));attrs=dict(carrier.attrib);group.remove(carrier)
        for n,line in enumerate(lines):
            text=element('text');text.attrib.update(attrs);text.set('x',str(cx));text.set('y',str(start+n*size*1.45));text.set('data-pptx-text-width',str(w));text.text=line;group.append(text)
    return root

def preview_svg(page,course,number):
    return ET.tostring(resolved_svg(page,course,number),encoding='unicode')

def export_pptx(pages,course):
    if not pages:raise ValueError('课件尚未生成')
    presentation=Presentation();presentation.slide_width=Inches(40/3);presentation.slide_height=Inches(7.5)
    def coord(v):return Inches(float(v)/96)
    def color(value):return RGBColor.from_string(value.lstrip('#'))
    def flat(shape):
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        for effect in shape._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):
            effect.set('idx','0')
    for number,page in enumerate(pages,1):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        for node in resolved_svg(page,course,number).iter():
            kind=tag(node);shape=None
            if kind in ['rect','ellipse','circle']:
                if kind=='rect':x,y,w,h=[float(node.get(a,'0')) for a in ['x','y','width','height']]
                else:
                    rx=float(node.get('rx',node.get('r','0')));ry=float(node.get('ry',node.get('r','0')));x=float(node.get('cx','0'))-rx;y=float(node.get('cy','0'))-ry;w=rx*2;h=ry*2
                shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE if kind=='rect' else MSO_SHAPE.OVAL,coord(x),coord(y),coord(w),coord(h))
                if node.get('fill','none')!='none':shape.fill.solid();shape.fill.fore_color.rgb=color(node.get('fill'))
                else:shape.fill.background()
            elif kind=='line':
                shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,*[coord(node.get(a,'0')) for a in ['x1','y1','x2','y2']])
            elif kind=='text':
                size=float(node.get('font-size','28'));value=node.text or '';anchor=node.get('text-anchor','start');w=float(node.get('data-pptx-text-width',max(size,units(value)*size+12)));x=float(node.get('x','0'))
                if anchor=='middle':x-=w/2
                if anchor=='end':x-=w
                box=slide.shapes.add_textbox(coord(x),coord(float(node.get('y','0'))-size*1.05),coord(w),coord(size*1.5))
                flat(box)
                frame=box.text_frame;frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.NONE;frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
                p=frame.paragraphs[0];p.space_before=p.space_after=Pt(0);p.line_spacing=1.0;p.alignment=PP_ALIGN.CENTER if anchor=='middle' else PP_ALIGN.RIGHT if anchor=='end' else PP_ALIGN.LEFT
                run=p.add_run();run.text=value;run.font.name='Arial';run.font.size=Pt(size*.75);run.font.bold=node.get('font-weight')=='bold';run.font.color.rgb=color(node.get('fill','#273B53'))
                ea=OxmlElement('a:ea');ea.set('typeface','Microsoft YaHei');run._r.get_or_add_rPr().append(ea)
            if shape is not None:
                flat(shape)
                if node.get('stroke'):
                    shape.line.color.rgb=color(node.get('stroke'));shape.line.width=Pt(float(node.get('stroke-width','1'))*.75)
                else:shape.line.fill.background()
        slide.notes_slide.notes_text_frame.text=page['content']
    output=BytesIO();presentation.save(output);return output.getvalue()
