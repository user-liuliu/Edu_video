"""Bounded document parsing. Uploaded content never supplies executable instructions."""
import io
import os
import signal
import subprocess
import tempfile
import zipfile
from pathlib import Path, PurePosixPath

import pymupdf as fitz
from PIL import Image

MAX_EXPANDED = 200 * 1024 * 1024
MAX_ITEMS = 5000
MAX_PAGES = 300
MIME = {'.pdf':'application/pdf', '.pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation',
        '.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        '.ppt':'application/vnd.ms-powerpoint', '.doc':'application/msword',
        '.png':'image/png', '.jpg':'image/jpeg', '.jpeg':'image/jpeg', '.webp':'image/webp'}


def supported_image_blob(data):
    """Return raster image bytes that the vision model can consume, else None."""
    try:
        with Image.open(io.BytesIO(data)) as image:
            if image.format not in {'PNG', 'JPEG', 'WEBP'}:
                return None
            if image.width * image.height > 40_000_000 or getattr(image, 'n_frames', 1) != 1:
                return None
            image.verify()
        return data
    except Exception:
        return None


def validate_document(filename, data):
    ext = Path(filename).suffix.lower()
    if ext not in MIME or not data:
        raise ValueError('不支持该文件类型或文件内容为空')
    try:
        if ext in ('.pptx', '.docx'):
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                items = archive.infolist()
                if len(items) > MAX_ITEMS or sum(i.file_size for i in items) > MAX_EXPANDED:
                    raise ValueError('压缩文档内容超出解析限制')
                for item in items:
                    path = PurePosixPath(item.filename.replace('\\', '/'))
                    if path.is_absolute() or '..' in path.parts or ':' in item.filename or item.flag_bits & 1:
                        raise ValueError('文档包含不安全路径或加密内容')
                    if 'vbaproject' in item.filename.lower() or (item.external_attr >> 16) & 0o170000 == 0o120000:
                        raise ValueError('不支持含宏或符号链接的文档')
                required = 'word/document.xml' if ext == '.docx' else 'ppt/presentation.xml'
                if required not in archive.namelist() or '[Content_Types].xml' not in archive.namelist():
                    raise ValueError('文档实际类型与扩展名不符')
        elif ext == '.pdf':
            if not data.startswith(b'%PDF-'):
                raise ValueError('PDF 文件格式不正确')
            with fitz.open(stream=data, filetype='pdf') as doc:
                if doc.needs_pass:
                    raise ValueError('不支持密码保护的文档')
                if not 0 < len(doc) <= MAX_PAGES:
                    raise ValueError('文档页数需在 1—300 页之间')
        elif ext in ('.doc', '.ppt'):
            if not data.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
                raise ValueError('旧版 Office 文件格式不正确')
            if b'EncryptedPackage' in data or 'EncryptedPackage'.encode('utf-16le') in data:
                raise ValueError('不支持密码保护的文档')
            import olefile
            with olefile.OleFileIO(io.BytesIO(data)) as legacy:
                entries=['/'.join(p).lower() for p in legacy.listdir()]
                if any('vba' in p or 'macros' in p or 'encrypted' in p for p in entries):
                    raise ValueError('不支持带宏或加密的旧版文档')
                required='WordDocument' if ext=='.doc' else 'PowerPoint Document'
                if not legacy.exists(required): raise ValueError('文档实际类型与扩展名不符')
        else:
            with Image.open(io.BytesIO(data)) as img:
                expected = {'.png':'PNG','.jpg':'JPEG','.jpeg':'JPEG','.webp':'WEBP'}[ext]
                if img.format != expected or img.width * img.height > 40_000_000 or getattr(img, 'n_frames', 1) != 1:
                    raise ValueError('图片格式不符、含动画或像素超过 4000 万')
                img.verify()
    except ValueError:
        raise
    except Exception:
        raise ValueError('文件已损坏、被加密或实际类型与扩展名不符') from None
    return MIME[ext]


def _restrict_converter():
    import ctypes
    import resource
    resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))
    resource.setrlimit(resource.RLIMIT_CPU, (60, 60))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_EXPANDED, MAX_EXPANDED))
    lib = ctypes.CDLL('libseccomp.so.2', use_errno=True)
    lib.seccomp_init.restype = ctypes.c_void_p
    lib.seccomp_init.argtypes = [ctypes.c_uint32]
    lib.seccomp_rule_add.argtypes = [ctypes.c_void_p, ctypes.c_uint32, ctypes.c_int, ctypes.c_uint]
    lib.seccomp_load.argtypes = [ctypes.c_void_p]
    lib.seccomp_syscall_resolve_name.argtypes = [ctypes.c_char_p]
    context = lib.seccomp_init(0x7fff0000)
    if not context:
        raise RuntimeError('无法隔离转换进程')
    for syscall in (b'connect', b'sendto', b'sendmsg'):
        if lib.seccomp_rule_add(context, 0x00050001, lib.seccomp_syscall_resolve_name(syscall), 0):
            raise RuntimeError('无法限制转换网络')
    if lib.seccomp_load(context):
        raise RuntimeError('无法启用转换隔离')


def convert_legacy(filename, data, temp_dir=None):
    if os.name != 'posix':
        raise ValueError('旧版 PPT/DOC 请在部署环境解析，或先另存为 PPTX/DOCX')
    target = 'pptx' if Path(filename).suffix.lower() == '.ppt' else 'docx'
    with tempfile.TemporaryDirectory(dir=temp_dir) as folder:
        root = Path(folder)
        source = root / ('input' + Path(filename).suffix.lower())
        source.write_bytes(data)
        profile = root / 'profile' / 'user'
        profile.mkdir(parents=True)
        (profile/'registrymodifications.xcu').write_text(
            '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
            '</oor:items>', encoding='utf-8')
        try:
            process = subprocess.Popen(['libreoffice', '-env:UserInstallation='+(root/'profile').as_uri(),
                '--headless', '--convert-to', target, '--outdir', str(root), str(source)],
                stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL,preexec_fn=_restrict_converter,start_new_session=True)
            try: process.wait(timeout=70)
            except subprocess.TimeoutExpired:
                os.killpg(process.pid,signal.SIGKILL);process.wait()
                raise
        except (subprocess.TimeoutExpired, FileNotFoundError, subprocess.SubprocessError):
            raise ValueError('旧版文档转换超时或转换服务不可用，请另存为现代格式后重试') from None
        output = root / ('input.' + target)
        if process.returncode or not output.is_file() or output.stat().st_size > MAX_EXPANDED:
            raise ValueError('文档转换失败，请检查文件是否损坏或加密')
        converted = output.read_bytes()
        validate_document(output.name, converted)
        return output.name, converted


def parse_document(filename, data, temp_dir=None, *, warnings=None):
    validate_document(filename, data)
    ext = Path(filename).suffix.lower()
    if ext in ('.doc', '.ppt'):
        filename, data = convert_legacy(filename, data, temp_dir)
        ext = Path(filename).suffix.lower()
    chunks = []
    def warn_image(location):
        if warnings is not None:
            warnings.append({'location': location, 'message': '该图片格式不受支持、已损坏或超过识别限制，未解析。请将重要公式、图表另存为 PNG/JPEG/WebP 后补传。'})
    def add(location, text='', image=None):
        parts=[text[n:n+6000] for n in range(0,len(text),6000)] or ['']
        for n,part in enumerate(parts):
            label=location if len(parts)==1 else f'{location} · 切片 {n+1}'
            chunks.append({'location':label, 'text':part, **({'image':image} if image and n==0 else {})})
        if len(chunks) > 1000:
            raise ValueError('素材切片超过 1000 个，请拆分文件')
    if ext == '.pdf':
        with fitz.open(stream=data, filetype='pdf') as doc:
            for n, page in enumerate(doc):
                txt = page.get_text().strip()
                if txt:
                    add(f'第 {n+1} 页', txt)
                if not txt or page.get_images():
                    scale=min(1,1800/max(page.rect.width,page.rect.height))
                    add(f'第 {n+1} 页（图像）', image=page.get_pixmap(matrix=fitz.Matrix(scale,scale)).tobytes('png'))
    elif ext == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        for n, p in enumerate(doc.paragraphs):
            if p.text.strip(): add(f'第 {n+1} 段', p.text.strip())
        for n, table in enumerate(doc.tables):
            add(f'第 {n+1} 个表格', '\n'.join(' | '.join(c.text for c in row.cells) for row in table.rows))
        for n, rel in enumerate(doc.part.rels.values()):
            if not rel.is_external and 'image' in rel.reltype:
                image = supported_image_blob(rel.target_part.blob)
                if image:
                    add(f'文档插图 {n+1}', image=image)
                else:
                    warn_image(f'文档插图 {n+1}')
    elif ext == '.pptx':
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > MAX_PAGES:
            raise ValueError('幻灯片不能超过 300 页')
        for n, slide in enumerate(presentation.slides):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame: parts.append(shape.text)
                if shape.has_table: parts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                if shape.shape_type == 13:
                    image = supported_image_blob(shape.image.blob)
                    if image: add(f'第 {n+1} 张幻灯片（插图）', image=image)
                    else: warn_image(f'第 {n+1} 张幻灯片（插图）')
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
            add(f'第 {n+1} 张幻灯片', '\n'.join(parts).strip())
    else:
        with Image.open(io.BytesIO(data)) as img:
            img.thumbnail((1800,1800))
            stream = io.BytesIO(); img.convert('RGB').save(stream, format='PNG')
            add('图片 1', image=stream.getvalue())
    if not chunks or not any(c['text'] or c.get('image') for c in chunks):
        raise ValueError('未提取到可用内容，请检查文档')
    return chunks
